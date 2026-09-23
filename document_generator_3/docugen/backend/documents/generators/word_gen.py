"""Générateur Word (python-docx) — document d'analyse."""
import io

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .base import GenerationContext, resolve_value

PRIMARY = RGBColor(0x25, 0x63, 0xEB)
DARK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)


def _set_cell_bg(cell, hex_color):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def generate(ctx: GenerationContext) -> bytes:
    doc = DocxDocument()

    # Bandeau entreprise + confidentialité
    header = doc.add_paragraph()
    run = header.add_run(ctx.company.name or "Entreprise")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = PRIMARY
    conf = header.add_run(f"    [{ctx.confidentiality_label.upper()}]")
    conf.bold = True
    conf.font.size = Pt(9)
    conf.font.color.rgb = MUTED

    # Titre
    title = doc.add_heading(level=0)
    trun = title.add_run(resolve_value(ctx, "title", ctx.document.title))
    trun.font.color.rgb = DARK
    sub = doc.add_paragraph()
    srun = sub.add_run(
        f"{resolve_value(ctx, 'client_name')} — "
        f"{resolve_value(ctx, 'project_name')}  ·  v{ctx.version_number}")
    srun.font.color.rgb = MUTED
    srun.italic = True

    # Tableau méta
    meta = [
        ("Projet", resolve_value(ctx, "project_name")),
        ("Client", resolve_value(ctx, "client_name")),
        ("Référence", resolve_value(ctx, "project_reference") or "—"),
        ("Confidentialité", ctx.confidentiality_label),
        ("Version", f"v{ctx.version_number}"),
        ("Auteur", f"{ctx.author_name} ({ctx.author_initials})".strip()),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in meta:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = str(value)
        row[0].paragraphs[0].runs[0].bold = True

    doc.add_paragraph()

    # Résumé exécutif
    summary = resolve_value(ctx, "summary")
    if summary:
        doc.add_heading("Résumé exécutif", level=1)
        doc.add_paragraph(str(summary))

    # Sections d'analyse : variable "sections" = [{"heading","content"}]
    sections = ctx.data.get("sections")
    if isinstance(sections, list):
        for sec in sections:
            if isinstance(sec, dict):
                if sec.get("heading"):
                    doc.add_heading(str(sec["heading"]), level=1)
                if sec.get("content"):
                    doc.add_paragraph(str(sec["content"]))
            else:
                doc.add_paragraph(str(sec))

    # Recommandations : variable "recommendations" = [str, ...]
    recs = ctx.data.get("recommendations")
    if isinstance(recs, list) and recs:
        doc.add_heading("Recommandations", level=1)
        for rec in recs:
            doc.add_paragraph(str(rec), style="List Bullet")

    # Historique des versions
    doc.add_page_break()
    doc.add_heading("Historique des versions", level=1)
    from documents.models import ConfidentialityLevel
    conf_map = dict(ConfidentialityLevel.choices)
    htable = doc.add_table(rows=1, cols=5)
    htable.style = "Light List Accent 1"
    hdr = htable.rows[0].cells
    for i, h in enumerate(["Version", "Date", "Auteur", "Confidentialité",
                           "Commentaire"]):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].bold = True
    versions = list(ctx.document.versions.order_by("version_number"))
    if not versions:
        cells = htable.add_row().cells
        cells[0].text = f"v{ctx.version_number}"
        cells[1].text = "en cours"
        cells[2].text = ctx.author_initials
        cells[3].text = ctx.confidentiality_label
        cells[4].text = ctx.comment or "—"
    for v in versions:
        cells = htable.add_row().cells
        cells[0].text = f"v{v.version_number}"
        cells[1].text = v.created_at.strftime("%d/%m/%Y %H:%M")
        cells[2].text = v.author_initials
        cells[3].text = conf_map.get(v.confidentiality, v.confidentiality)
        cells[4].text = v.comment or "—"

    # Pied entreprise
    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bits = [b for b in [ctx.company.website_url, ctx.company.email,
                        ctx.company.phone] if b]
    frun = foot.add_run("  ·  ".join(bits))
    frun.font.size = Pt(8)
    frun.font.color.rgb = MUTED

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
