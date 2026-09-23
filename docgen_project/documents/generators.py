"""Générateurs de documents.

Chaque générateur reçoit un objet `Document` et retourne un tuple
(contenu_binaire, nom_de_fichier). Les données/variables du document
(`document.data`) alimentent le contenu, et le profil entreprise ainsi que
le projet fournissent l'en-tête, le pied de page et le bandeau de
confidentialité.
"""
import io
from datetime import datetime

from django.utils import timezone

from core.models import CompanyProfile


CONF_LABELS = {
    "public": "PUBLIC",
    "internal": "INTERNE",
    "confidential": "CONFIDENTIEL",
    "secret": "SECRET",
}


def build_context(document):
    """Assemble un dictionnaire de contexte commun à tous les générateurs."""
    company = CompanyProfile.get_solo()
    project = document.project
    return {
        "company": company,
        "project": project,
        "document": document,
        "data": document.data or {},
        "confidentiality": CONF_LABELS.get(document.confidentiality, "INTERNE"),
        "generated_at": timezone.localtime().strftime("%d/%m/%Y %H:%M"),
    }


# =========================================================================
# PDF — Suivi de projet (reportlab / Platypus)
# =========================================================================

def generate_pdf_tracking(document):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image,
    )
    from reportlab.lib.enums import TA_CENTER

    ctx = build_context(document)
    company, project, data = ctx["company"], ctx["project"], ctx["data"]

    buffer = io.BytesIO()
    conf = ctx["confidentiality"]

    def header_footer(canvas, doc):
        canvas.saveState()
        # Bandeau de confidentialité (haut)
        canvas.setFillColor(colors.HexColor("#c62828"))
        canvas.setFont("Helvetica-Bold", 8)
        canvas.drawCentredString(A4[0] / 2, A4[1] - 1.0 * cm, f"— {conf} —")
        # Pied de page
        canvas.setFillColor(colors.grey)
        canvas.setFont("Helvetica", 7)
        footer = f"{company.name} · {company.website or ''} · {company.email or ''}"
        canvas.drawCentredString(A4[0] / 2, 1.1 * cm, footer.strip(" ·"))
        canvas.drawRightString(A4[0] - 2 * cm, 1.1 * cm, f"Page {doc.page}")
        canvas.drawString(2 * cm, 1.1 * cm, f"{conf}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        topMargin=1.6 * cm, bottomMargin=1.8 * cm,
        leftMargin=2 * cm, rightMargin=2 * cm,
        title=document.title, author=company.name,
    )

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Small", parent=styles["Normal"], fontSize=8, textColor=colors.grey))
    styles.add(ParagraphStyle(name="H1c", parent=styles["Heading1"], textColor=colors.HexColor("#1565c0")))
    styles.add(ParagraphStyle(name="Center", parent=styles["Normal"], alignment=TA_CENTER))

    story = []

    # En-tête : logo entreprise + coordonnées
    header_cells = []
    logo_flowable = ""
    if company.logo and hasattr(company.logo, "path"):
        try:
            logo_flowable = Image(company.logo.path, width=3.5 * cm, height=1.8 * cm, kind="proportional")
        except Exception:
            logo_flowable = Paragraph(f"<b>{company.name}</b>", styles["Heading2"])
    else:
        logo_flowable = Paragraph(f"<b>{company.name}</b>", styles["Heading2"])

    company_info = Paragraph(
        f"<b>{company.name}</b><br/>{(company.address or '').replace(chr(10), '<br/>')}<br/>"
        f"{company.phone or ''} · {company.email or ''}",
        styles["Small"],
    )
    header_table = Table([[logo_flowable, company_info]], colWidths=[6 * cm, 11 * cm])
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (1, 0), (1, 0), "RIGHT"),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 0.3 * cm))
    story.append(Table([[""]], colWidths=[17 * cm], style=TableStyle([
        ("LINEBELOW", (0, 0), (-1, -1), 1, colors.HexColor("#1565c0")),
    ])))
    story.append(Spacer(1, 0.5 * cm))

    # Titre
    story.append(Paragraph(document.title, styles["H1c"]))
    story.append(Paragraph(f"Suivi de projet · Généré le {ctx['generated_at']} · Version {document.current_version + 1}", styles["Small"]))
    story.append(Spacer(1, 0.5 * cm))

    # Bloc infos projet
    info_rows = [
        ["Projet", project.name],
        ["Client", project.client_name],
        ["Contact client", project.client_contact.full_name if project.client_contact else "—"],
        ["Contact interne", project.internal_contact.full_name if project.internal_contact else "—"],
        ["Confidentialité", conf],
    ]
    info_table = Table(info_rows, colWidths=[5 * cm, 12 * cm])
    info_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#e3f2fd")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#bbdefb")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(info_table)
    story.append(Spacer(1, 0.5 * cm))

    if project.description:
        story.append(Paragraph("<b>Description du projet</b>", styles["Heading3"]))
        story.append(Paragraph(project.description.replace("\n", "<br/>"), styles["Normal"]))
        story.append(Spacer(1, 0.5 * cm))

    # Rendu des variables/données
    _render_data_to_pdf(story, data, styles, colors, Table, TableStyle, Spacer, Paragraph, cm)

    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    buffer.seek(0)
    filename = f"{_slug(document.title)}_v{document.current_version + 1}.pdf"
    return buffer.getvalue(), filename


def _render_data_to_pdf(story, data, styles, colors, Table, TableStyle, Spacer, Paragraph, cm):
    """Convertit les variables (dict) en sections lisibles dans le PDF."""
    if not data:
        return
    story.append(Paragraph("<b>Données du document</b>", styles["Heading2"]))
    story.append(Spacer(1, 0.2 * cm))
    for key, value in data.items():
        label = str(key).replace("_", " ").capitalize()
        if isinstance(value, list):
            story.append(Paragraph(f"<b>{label}</b>", styles["Heading4"]))
            # Liste de dicts -> tableau ; liste simple -> puces
            if value and isinstance(value[0], dict):
                cols = list(value[0].keys())
                rows = [[str(c).capitalize() for c in cols]]
                for item in value:
                    rows.append([str(item.get(c, "")) for c in cols])
                t = Table(rows, hAlign="LEFT")
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1565c0")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
                ]))
                story.append(t)
            else:
                for item in value:
                    story.append(Paragraph(f"• {item}", styles["Normal"]))
            story.append(Spacer(1, 0.3 * cm))
        elif isinstance(value, dict):
            story.append(Paragraph(f"<b>{label}</b>", styles["Heading4"]))
            rows = [[str(k).capitalize(), str(v)] for k, v in value.items()]
            t = Table(rows, colWidths=[5 * cm, 12 * cm])
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ]))
            story.append(t)
            story.append(Spacer(1, 0.3 * cm))
        else:
            story.append(Paragraph(f"<b>{label} :</b> {value}", styles["Normal"]))
            story.append(Spacer(1, 0.15 * cm))


# =========================================================================
# EXCEL — générique (openpyxl)
# =========================================================================

def generate_excel(document):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    ctx = build_context(document)
    company, project, data = ctx["company"], ctx["project"], ctx["data"]

    wb = Workbook()

    # --- Feuille d'informations ---
    ws = wb.active
    ws.title = "Informations"
    title_font = Font(size=14, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1565C0")
    label_font = Font(bold=True)

    ws.merge_cells("A1:B1")
    ws["A1"] = document.title
    ws["A1"].font = title_font
    ws["A1"].fill = header_fill
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A2"] = f"Confidentialité : {ctx['confidentiality']}"
    ws["A2"].font = Font(bold=True, color="C62828")

    rows = [
        ("Entreprise", company.name),
        ("Projet", project.name),
        ("Client", project.client_name),
        ("Contact client", project.client_contact.full_name if project.client_contact else "—"),
        ("Contact interne", project.internal_contact.full_name if project.internal_contact else "—"),
        ("Généré le", ctx["generated_at"]),
        ("Version", document.current_version + 1),
    ]
    r = 4
    for label, value in rows:
        ws.cell(row=r, column=1, value=label).font = label_font
        ws.cell(row=r, column=2, value=str(value))
        r += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 45

    # --- Données ---
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for key, value in (data or {}).items():
        sheet_name = str(key)[:28] or "Donnees"
        # Éviter les doublons de nom de feuille
        base = sheet_name
        i = 1
        while sheet_name in wb.sheetnames:
            sheet_name = f"{base[:25]}_{i}"
            i += 1
        wsd = wb.create_sheet(title=sheet_name)

        if isinstance(value, list) and value and isinstance(value[0], dict):
            cols = list(value[0].keys())
            for c, col in enumerate(cols, start=1):
                cell = wsd.cell(row=1, column=c, value=str(col).capitalize())
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.border = border
            for rr, item in enumerate(value, start=2):
                for c, col in enumerate(cols, start=1):
                    cell = wsd.cell(row=rr, column=c, value=item.get(col, ""))
                    cell.border = border
            for c in range(1, len(cols) + 1):
                wsd.column_dimensions[chr(64 + c)].width = 20
        elif isinstance(value, list):
            wsd.cell(row=1, column=1, value=str(key).capitalize()).font = Font(bold=True)
            for rr, item in enumerate(value, start=2):
                wsd.cell(row=rr, column=1, value=str(item))
            wsd.column_dimensions["A"].width = 40
        elif isinstance(value, dict):
            wsd.cell(row=1, column=1, value="Clé").font = Font(bold=True)
            wsd.cell(row=1, column=2, value="Valeur").font = Font(bold=True)
            for rr, (k, v) in enumerate(value.items(), start=2):
                wsd.cell(row=rr, column=1, value=str(k))
                wsd.cell(row=rr, column=2, value=str(v))
            wsd.column_dimensions["A"].width = 25
            wsd.column_dimensions["B"].width = 40
        else:
            wsd.cell(row=1, column=1, value=str(key).capitalize()).font = Font(bold=True)
            wsd.cell(row=2, column=1, value=str(value))
            wsd.column_dimensions["A"].width = 40

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"{_slug(document.title)}_v{document.current_version + 1}.xlsx"
    return buffer.getvalue(), filename


# =========================================================================
# WORD — analyse (python-docx)
# =========================================================================

def generate_word_analysis(document):
    from docx import Document as DocxDocument
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    ctx = build_context(document)
    company, project, data = ctx["company"], ctx["project"], ctx["data"]

    docx = DocxDocument()

    # Bandeau de confidentialité en en-tête
    section = docx.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.text = f"— {ctx['confidentiality']} —"
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in hp.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
        run.font.bold = True

    # Pied de page
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = f"{company.name} · {company.website or ''} · {company.email or ''}".strip(" ·")
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in fp.runs:
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    # Logo entreprise
    if company.logo and hasattr(company.logo, "path"):
        try:
            docx.add_picture(company.logo.path, width=Cm(4))
        except Exception:
            pass

    # Titre
    h = docx.add_heading(document.title, level=0)
    sub = docx.add_paragraph(
        f"Document d'analyse · {company.name} · Généré le {ctx['generated_at']} · Version {document.current_version + 1}"
    )
    sub.runs[0].font.size = Pt(9)
    sub.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    # Tableau d'informations
    docx.add_heading("Informations générales", level=1)
    info = [
        ("Projet", project.name),
        ("Client", project.client_name),
        ("Contact client", project.client_contact.full_name if project.client_contact else "—"),
        ("Contact interne", project.internal_contact.full_name if project.internal_contact else "—"),
        ("Confidentialité", ctx["confidentiality"]),
    ]
    table = docx.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in info:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)
        cells[0].paragraphs[0].runs[0].font.bold = True

    if project.description:
        docx.add_heading("Description du projet", level=1)
        docx.add_paragraph(project.description)

    # Données/variables
    if data:
        docx.add_heading("Analyse détaillée", level=1)
        _render_data_to_docx(docx, data)

    buffer = io.BytesIO()
    docx.save(buffer)
    buffer.seek(0)
    filename = f"{_slug(document.title)}_v{document.current_version + 1}.docx"
    return buffer.getvalue(), filename


def _render_data_to_docx(docx, data):
    for key, value in data.items():
        label = str(key).replace("_", " ").capitalize()
        if isinstance(value, list) and value and isinstance(value[0], dict):
            docx.add_heading(label, level=2)
            cols = list(value[0].keys())
            table = docx.add_table(rows=1, cols=len(cols))
            table.style = "Light List Accent 1"
            for c, col in enumerate(cols):
                table.rows[0].cells[c].text = str(col).capitalize()
            for item in value:
                cells = table.add_row().cells
                for c, col in enumerate(cols):
                    cells[c].text = str(item.get(col, ""))
        elif isinstance(value, list):
            docx.add_heading(label, level=2)
            for item in value:
                docx.add_paragraph(str(item), style="List Bullet")
        elif isinstance(value, dict):
            docx.add_heading(label, level=2)
            table = docx.add_table(rows=0, cols=2)
            table.style = "Light Grid Accent 1"
            for k, v in value.items():
                cells = table.add_row().cells
                cells[0].text = str(k)
                cells[1].text = str(v)
        else:
            p = docx.add_paragraph()
            run = p.add_run(f"{label} : ")
            run.bold = True
            p.add_run(str(value))


# =========================================================================
# Dispatcher
# =========================================================================

GENERATORS = {
    "pdf_tracking": generate_pdf_tracking,
    "excel": generate_excel,
    "word_analysis": generate_word_analysis,
}


def generate_document(document):
    """Point d'entrée : retourne (contenu_binaire, nom_de_fichier)."""
    generator = GENERATORS.get(document.doc_type)
    if not generator:
        raise ValueError(f"Type de document non supporté : {document.doc_type}")
    return generator(document)


def _slug(text):
    import re
    text = (text or "document").lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_") or "document"
