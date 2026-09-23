"""
Générateur « par blocs » : rend un document à partir de la liste de blocs
définie visuellement dans le modèle.

Types de blocs pris en charge :
  - heading    : titre fixe (niveau 1-3), interpolation {{variables}}
  - paragraph  : paragraphe fixe, interpolation {{variables}}
  - field      : champ rempli par l'utilisateur (texte, date, liste…)
  - table      : tableau (colonnes + lignes) rempli façon tableur
  - contacts   : insertion des contacts du projet (client / interne / les deux)
  - spacer     : espacement
"""
import io

from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, Spacer, Table, TableStyle

from .base import GenerationContext, interpolate, placeholder_context
from . import pdf_gen as P


# ---------------------------------------------------------------------------
# Helpers communs
# ---------------------------------------------------------------------------
def _field_display(block, ctx):
    """Valeur d'un champ, éventuellement préfixée par son libellé."""
    key = block.get("key", "")
    value = ctx.data.get(key, "")
    if isinstance(value, (list, dict)):
        value = ""
    value = "" if value is None else str(value)
    label = block.get("label", "")
    if block.get("show_label", True) and label:
        return label, value
    return None, value


def _table_data(block, ctx):
    """Retourne (colonnes, lignes) pour un bloc tableau."""
    key = block.get("key", "")
    stored = ctx.data.get(key) or {}
    columns = stored.get("columns")
    if not columns:
        columns = [c.get("label", "") for c in block.get("columns", [])]
    rows = stored.get("rows", [])
    # Normalise la longueur des lignes
    width = len(columns)
    norm = []
    for r in rows:
        r = list(r) + [""] * (width - len(r))
        norm.append([("" if c is None else str(c)) for c in r[:width]])
    return columns, norm


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def render_pdf(ctx: GenerationContext) -> bytes:
    styles = P._styles()
    pctx = placeholder_context(ctx)
    buffer = io.BytesIO()
    doc = P._DocTemplate(buffer, ctx)
    story = []

    for block in ctx.document.template.schema:
        btype = block.get("type")

        if btype == "heading":
            level = int(block.get("level", 2))
            style = styles["DGTitle"] if level == 1 else styles["DGH2"]
            story.append(Paragraph(interpolate(block.get("text", ""), pctx), style))

        elif btype == "paragraph":
            txt = interpolate(block.get("text", ""), pctx).replace("\n", "<br/>")
            story.append(Paragraph(txt, styles["DGBody"]))

        elif btype == "field":
            label, value = _field_display(block, ctx)
            value = value.replace("\n", "<br/>")
            if label:
                story.append(Paragraph(f"<b>{label} :</b> {value}", styles["DGBody"]))
            else:
                story.append(Paragraph(value, styles["DGBody"]))

        elif btype == "table":
            columns, rows = _table_data(block, ctx)
            if block.get("label"):
                story.append(Paragraph(block["label"], styles["DGH2"]))
            if columns:
                data = [columns] + (rows or [[""] * len(columns)])
                t = Table(data, hAlign="LEFT")
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), P.PRIMARY),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, P.LIGHT]),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E2E8F0")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ]))
                story.append(t)

        elif btype == "contacts":
            story += _contacts_pdf(block, ctx, styles)

        elif btype == "spacer":
            story.append(Spacer(1, 12))

    # Historique + pied entreprise (toujours ajoutés)
    story.append(PageBreak())
    story += P._history_block(ctx, styles)
    story.append(Spacer(1, 20))
    story += P._company_footer_block(ctx, styles)

    doc.build(story)
    return buffer.getvalue()


def _contacts_pdf(block, ctx, styles):
    scope = block.get("scope", "both")
    contacts = []
    if scope in ("client", "both"):
        contacts += [("Client", c) for c in ctx.client_contacts()]
    if scope in ("internal", "both"):
        contacts += [("Interne", c) for c in ctx.internal_contacts()]
    out = [Paragraph(block.get("label", "Contacts"), styles["DGH2"])]
    if not contacts:
        out.append(Paragraph("Aucun contact renseigné.", styles["DGBody"]))
        return out
    header = [["Rôle", "Nom", "Fonction", "Email", "Téléphone"]]
    body = [[kind, c.full_name, c.role or "—", c.email or "—", c.phone or "—"]
            for kind, c in contacts]
    t = Table(header + body, colWidths=[2 * cm, 3.5 * cm, 3 * cm, 4 * cm, 3 * cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), P.PRIMARY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, P.LIGHT]),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E2E8F0")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    out.append(t)
    return out


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def render_docx(ctx: GenerationContext) -> bytes:
    from docx import Document as DocxDocument
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    PRIMARY = RGBColor(0x25, 0x63, 0xEB)
    MUTED = RGBColor(0x64, 0x74, 0x8B)
    pctx = placeholder_context(ctx)

    doc = DocxDocument()

    # Bandeau entreprise + confidentialité
    header = doc.add_paragraph()
    run = header.add_run(ctx.company.name or "Entreprise")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = PRIMARY
    conf = header.add_run(f"    [{ctx.confidentiality_label.upper()}]")
    conf.bold = True
    conf.font.size = Pt(9)
    conf.font.color.rgb = MUTED

    for block in ctx.document.template.schema:
        btype = block.get("type")

        if btype == "heading":
            level = int(block.get("level", 2))
            doc.add_heading(interpolate(block.get("text", ""), pctx),
                            level=0 if level == 1 else min(level, 3))

        elif btype == "paragraph":
            doc.add_paragraph(interpolate(block.get("text", ""), pctx))

        elif btype == "field":
            label, value = _field_display(block, ctx)
            p = doc.add_paragraph()
            if label:
                r = p.add_run(f"{label} : ")
                r.bold = True
            p.add_run(value)

        elif btype == "table":
            columns, rows = _table_data(block, ctx)
            if block.get("label"):
                doc.add_heading(block["label"], level=2)
            if columns:
                table = doc.add_table(rows=1, cols=len(columns))
                table.style = "Light Grid Accent 1"
                for i, col in enumerate(columns):
                    cell = table.rows[0].cells[i]
                    cell.text = str(col)
                    if cell.paragraphs[0].runs:
                        cell.paragraphs[0].runs[0].bold = True
                for r in (rows or []):
                    cells = table.add_row().cells
                    for i, val in enumerate(r):
                        cells[i].text = str(val)

        elif btype == "contacts":
            _contacts_docx(doc, block, ctx)

        elif btype == "spacer":
            doc.add_paragraph()

    # Historique
    doc.add_page_break()
    doc.add_heading("Historique des versions", level=1)
    from documents.models import ConfidentialityLevel
    conf_map = dict(ConfidentialityLevel.choices)
    htable = doc.add_table(rows=1, cols=5)
    htable.style = "Light List Accent 1"
    for i, h in enumerate(["Version", "Date", "Auteur", "Confidentialité",
                           "Commentaire"]):
        htable.rows[0].cells[i].text = h
        if htable.rows[0].cells[i].paragraphs[0].runs:
            htable.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    versions = list(ctx.document.versions.order_by("version_number"))
    if not versions:
        cells = htable.add_row().cells
        cells[0].text = f"v{ctx.version_number}"
        cells[1].text = "en cours"
        cells[2].text = ctx.author_initials
        cells[3].text = ctx.confidentiality_label
        cells[4].text = ctx.comment or "—"
    for v in versions:
        cells = htable.add_row().cells
        cells[0].text = f"v{v.version_number}"
        cells[1].text = v.created_at.strftime("%d/%m/%Y %H:%M")
        cells[2].text = v.author_initials
        cells[3].text = conf_map.get(v.confidentiality, v.confidentiality)
        cells[4].text = v.comment or "—"

    # Pied
    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bits = [b for b in [ctx.company.website_url, ctx.company.email,
                        ctx.company.phone] if b]
    frun = foot.add_run("  ·  ".join(bits))
    frun.font.size = Pt(8)
    frun.font.color.rgb = MUTED

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _contacts_docx(doc, block, ctx):
    scope = block.get("scope", "both")
    contacts = []
    if scope in ("client", "both"):
        contacts += [("Client", c) for c in ctx.client_contacts()]
    if scope in ("internal", "both"):
        contacts += [("Interne", c) for c in ctx.internal_contacts()]
    doc.add_heading(block.get("label", "Contacts"), level=2)
    if not contacts:
        doc.add_paragraph("Aucun contact renseigné.")
        return
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light List Accent 1"
    for i, h in enumerate(["Rôle", "Nom", "Fonction", "Email", "Téléphone"]):
        table.rows[0].cells[i].text = h
        if table.rows[0].cells[i].paragraphs[0].runs:
            table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for kind, c in contacts:
        cells = table.add_row().cells
        cells[0].text = kind
        cells[1].text = c.full_name
        cells[2].text = c.role or "—"
        cells[3].text = c.email or "—"
        cells[4].text = c.phone or "—"


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def render_xlsx(ctx: GenerationContext) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HEADER_FILL = PatternFill("solid", fgColor="2563EB")
    ALT_FILL = PatternFill("solid", fgColor="F1F5F9")
    HEADER_FONT = Font(bold=True, color="FFFFFF")
    LABEL_FONT = Font(bold=True, color="1E293B")
    TITLE_FONT = Font(size=15, bold=True, color="1E293B")
    THIN = Side(style="thin", color="E2E8F0")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    pctx = placeholder_context(ctx)

    wb = Workbook()
    ws = wb.active
    ws.title = "Document"
    row = 1

    def autosize(sheet):
        for col in sheet.columns:
            length, letter = 0, None
            for c in col:
                if letter is None:
                    letter = get_column_letter(c.column)
                if c.value is not None:
                    length = max(length, len(str(c.value)))
            if letter:
                sheet.column_dimensions[letter].width = min(length + 4, 60)

    for block in ctx.document.template.schema:
        btype = block.get("type")

        if btype == "heading":
            cell = ws.cell(row=row, column=1,
                           value=interpolate(block.get("text", ""), pctx))
            cell.font = TITLE_FONT if int(block.get("level", 2)) == 1 else LABEL_FONT
            row += 2

        elif btype == "paragraph":
            ws.cell(row=row, column=1,
                    value=interpolate(block.get("text", ""), pctx))
            row += 2

        elif btype == "field":
            label, value = _field_display(block, ctx)
            if label:
                ws.cell(row=row, column=1, value=label).font = LABEL_FONT
                ws.cell(row=row, column=2, value=value)
            else:
                ws.cell(row=row, column=1, value=value)
            row += 1

        elif btype == "table":
            columns, rows = _table_data(block, ctx)
            if block.get("label"):
                ws.cell(row=row, column=1, value=block["label"]).font = LABEL_FONT
                row += 1
            if columns:
                for j, col in enumerate(columns, start=1):
                    cell = ws.cell(row=row, column=j, value=str(col))
                    cell.fill = HEADER_FILL
                    cell.font = HEADER_FONT
                    cell.alignment = Alignment(horizontal="center")
                    cell.border = BORDER
                row += 1
                for r in (rows or []):
                    for j, val in enumerate(r, start=1):
                        cell = ws.cell(row=row, column=j, value=val)
                        cell.border = BORDER
                        if row % 2 == 0:
                            cell.fill = ALT_FILL
                    row += 1
            row += 1

        elif btype == "contacts":
            scope = block.get("scope", "both")
            contacts = []
            if scope in ("client", "both"):
                contacts += [("Client", c) for c in ctx.client_contacts()]
            if scope in ("internal", "both"):
                contacts += [("Interne", c) for c in ctx.internal_contacts()]
            ws.cell(row=row, column=1,
                    value=block.get("label", "Contacts")).font = LABEL_FONT
            row += 1
            for j, h in enumerate(["Rôle", "Nom", "Fonction", "Email",
                                   "Téléphone"], start=1):
                cell = ws.cell(row=row, column=j, value=h)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.border = BORDER
            row += 1
            for kind, c in contacts:
                for j, val in enumerate([kind, c.full_name, c.role, c.email,
                                         c.phone], start=1):
                    ws.cell(row=row, column=j, value=val).border = BORDER
                row += 1
            row += 1

        elif btype == "spacer":
            row += 1

    autosize(ws)

    # Feuille Historique
    hist = wb.create_sheet("Historique")
    from documents.models import ConfidentialityLevel
    conf_map = dict(ConfidentialityLevel.choices)
    for j, h in enumerate(["Version", "Date", "Auteur", "Confidentialité",
                           "Commentaire"], start=1):
        cell = hist.cell(row=1, column=j, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    r = 2
    for v in ctx.document.versions.order_by("version_number"):
        hist.append([f"v{v.version_number}",
                     v.created_at.strftime("%d/%m/%Y %H:%M"),
                     v.author_initials,
                     conf_map.get(v.confidentiality, v.confidentiality),
                     v.comment])
        r += 1
    if r == 2:
        hist.append([f"v{ctx.version_number}", "en cours", ctx.author_initials,
                     ctx.confidentiality_label, ctx.comment])
    autosize(hist)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


BLOCK_RENDERERS = {
    "pdf": render_pdf,
    "docx": render_docx,
    "xlsx": render_xlsx,
}
