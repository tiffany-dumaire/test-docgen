"""Générateur Excel (openpyxl)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .base import GenerationContext, resolve_value

HEADER_FILL = PatternFill("solid", fgColor="2563EB")
ALT_FILL = PatternFill("solid", fgColor="F1F5F9")
TITLE_FONT = Font(size=16, bold=True, color="1E293B")
LABEL_FONT = Font(bold=True, color="1E293B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="E2E8F0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _autosize(ws, max_width=60):
    for col in ws.columns:
        length = 0
        letter = None
        for cell in col:
            if letter is None:
                letter = get_column_letter(cell.column)
            if cell.value is not None:
                length = max(length, len(str(cell.value)))
        if letter:
            ws.column_dimensions[letter].width = min(length + 4, max_width)


def generate(ctx: GenerationContext) -> bytes:
    wb = Workbook()

    # --- Feuille "Informations" ---
    info = wb.active
    info.title = "Informations"
    info["A1"] = ctx.company.name or "Entreprise"
    info["A1"].font = TITLE_FONT
    info.append([])
    meta = [
        ("Titre", ctx.document.title),
        ("Projet", resolve_value(ctx, "project_name")),
        ("Client", resolve_value(ctx, "client_name")),
        ("Référence", resolve_value(ctx, "project_reference")),
        ("Confidentialité", ctx.confidentiality_label),
        ("Version", f"v{ctx.version_number}"),
        ("Auteur", f"{ctx.author_name} ({ctx.author_initials})".strip()),
    ]
    for label, value in meta:
        info.append([label, value])
        info[f"A{info.max_row}"].font = LABEL_FONT
    _autosize(info)

    # --- Feuille de données générique ---
    # Variable "table" : {"columns": [...], "rows": [[...], ...]}
    table = ctx.data.get("table")
    data_ws = wb.create_sheet("Données")
    if isinstance(table, dict) and table.get("columns"):
        columns = table["columns"]
        rows = table.get("rows", [])
        for idx, col in enumerate(columns, start=1):
            cell = data_ws.cell(row=1, column=idx, value=str(col))
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER
        for r_idx, row in enumerate(rows, start=2):
            for c_idx, val in enumerate(row, start=1):
                cell = data_ws.cell(row=r_idx, column=c_idx, value=val)
                cell.border = BORDER
                if r_idx % 2 == 0:
                    cell.fill = ALT_FILL
        data_ws.freeze_panes = "A2"
        _autosize(data_ws)
    else:
        data_ws["A1"] = "Aucune donnée fournie (variable 'table' absente)."

    # --- Feuille "Historique" ---
    hist = wb.create_sheet("Historique")
    headers = ["Version", "Date", "Auteur", "Confidentialité", "Commentaire"]
    for idx, h in enumerate(headers, start=1):
        cell = hist.cell(row=1, column=idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
    from documents.models import ConfidentialityLevel
    conf_map = dict(ConfidentialityLevel.choices)
    r = 2
    for v in ctx.document.versions.order_by("version_number"):
        hist.append([
            f"v{v.version_number}",
            v.created_at.strftime("%d/%m/%Y %H:%M"),
            v.author_initials,
            conf_map.get(v.confidentiality, v.confidentiality),
            v.comment,
        ])
        r += 1
    if r == 2:
        hist.append([f"v{ctx.version_number}", "en cours", ctx.author_initials,
                     ctx.confidentiality_label, ctx.comment])
    _autosize(hist)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
