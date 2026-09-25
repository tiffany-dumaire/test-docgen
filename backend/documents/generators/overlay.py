"""
Calques dynamiques (« overlays ») : placement libre d'éléments par-dessus un
document déjà généré.

Deux cibles :
  - PDF  : chaque calque est estampillé (pypdf) sur les pages visées d'un PDF
           déjà produit (page de garde comprise). Le rendu réutilise le moteur
           de mise en page libre (reportlab) sur un fond transparent.
  - PPTX : les éléments sont ajoutés comme *formes natives* (python-pptx) sur
           les diapositives visées, en conservant le template titre/contenu.

Un overlay stocké dans template.settings["overlays"] :
    {
      "id", "name", "enabled": true,
      "doc_type": "pdf|pptx",
      "target": {"pages": "all|first|last|odd|even|1,3|2-4"},
      "layout": {"page_size": "a4|a3|slide|slide43", "orientation", "elements": [...]}
    }
Les éléments partagent le format de layout_render (text/image/logo/rect/
ellipse/line, coordonnées en points, origine haut-gauche).
"""
import io

from .base import interpolate, placeholder_context
from . import layout_render as LR

EMU_PER_PT = 12700


# --------------------------------------------------------------------------- #
# Sélection des pages / diapos
# --------------------------------------------------------------------------- #
def parse_pages(spec, total: int):
    """Renvoie l'ensemble (trié) des index 0-based visés parmi `total` pages.

    Accepte : 'all', 'first', 'last', 'odd', 'even', ou une liste type
    '1,3,5' / '2-4' / '1,3-5,last' (numéros 1-based dans la spec).
    """
    if total <= 0:
        return []
    spec = str(spec or "all").strip().lower()
    if not spec or spec == "all":
        return list(range(total))
    if spec == "first":
        return [0]
    if spec == "last":
        return [total - 1]
    if spec == "odd":   # pages impaires (1,3,5…) -> index pairs
        return [i for i in range(total) if i % 2 == 0]
    if spec == "even":  # pages paires (2,4,6…) -> index impairs
        return [i for i in range(total) if i % 2 == 1]
    out = set()
    for part in spec.replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        if part == "first":
            out.add(0); continue
        if part == "last":
            out.add(total - 1); continue
        if "-" in part:
            a, _, b = part.partition("-")
            try:
                lo = 1 if a.strip() in ("", "first") else int(a)
                hi = total if b.strip() in ("", "last") else int(b)
            except ValueError:
                continue
            if lo > hi:
                lo, hi = hi, lo
            for n in range(lo, hi + 1):
                if 1 <= n <= total:
                    out.add(n - 1)
        else:
            try:
                n = int(part)
            except ValueError:
                continue
            if 1 <= n <= total:
                out.add(n - 1)
    return sorted(out)


def _overlays_for(document, doc_type):
    items = (document.template.settings or {}).get("overlays") or []
    out = []
    for ov in items:
        if not isinstance(ov, dict):
            continue
        if ov.get("enabled") is False:
            continue
        if (ov.get("doc_type") or doc_type) != doc_type:
            continue
        layout = ov.get("layout") or {}
        if not layout.get("elements"):
            continue
        out.append(ov)
    return out


def has_overlays(document, doc_type) -> bool:
    return bool(_overlays_for(document, doc_type))


# --------------------------------------------------------------------------- #
# PDF : estampillage via pypdf
# --------------------------------------------------------------------------- #
def _scaled_layout(layout, target_w, target_h):
    """Copie la mise en page en adaptant les coordonnées à (target_w, target_h)."""
    dw, dh = LR.page_dims(layout)
    if dw <= 0 or dh <= 0:
        return layout, 1.0
    sx = target_w / dw
    sy = target_h / dh
    if abs(sx - 1) < 1e-3 and abs(sy - 1) < 1e-3:
        return layout, 1.0
    fs = (sx + sy) / 2.0
    out = dict(layout)
    els = []
    for el in layout.get("elements", []):
        e = dict(el)
        for k, f in (("x", sx), ("y", sy), ("w", sx), ("h", sy)):
            if k in e:
                try:
                    e[k] = float(e[k]) * f
                except (TypeError, ValueError):
                    pass
        for k in ("size", "leading", "width", "stroke_width", "radius"):
            if k in e:
                try:
                    e[k] = float(e[k]) * fs
                except (TypeError, ValueError):
                    pass
        els.append(e)
    out["elements"] = els
    return out, fs


def _render_overlay_pdf_page(layout, ctx, page_w, page_h) -> bytes:
    """Page PDF transparente (sans fond) aux dimensions cibles avec les éléments."""
    from reportlab.pdfgen import canvas as pdfcanvas
    lay = dict(layout)
    lay.pop("background", None)  # transparent : on n'écrase pas la page de base
    scaled, _ = _scaled_layout(lay, page_w, page_h)
    buf = io.BytesIO()
    c = pdfcanvas.Canvas(buf, pagesize=(page_w, page_h))
    LR.draw_on_canvas(c, scaled, ctx, page_w, page_h)
    c.showPage()
    c.save()
    return buf.getvalue()


def stamp_pdf(base_pdf: bytes, document, ctx) -> bytes:
    """Estampille les calques PDF sur les pages visées d'un PDF déjà généré."""
    overlays = _overlays_for(document, "pdf")
    if not overlays:
        return base_pdf
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(io.BytesIO(base_pdf))
    total = len(reader.pages)
    if total == 0:
        return base_pdf

    # index de page -> liste de layouts à appliquer (dans l'ordre)
    plan = {}
    for ov in overlays:
        idxs = parse_pages((ov.get("target") or {}).get("pages", "all"), total)
        for i in idxs:
            plan.setdefault(i, []).append(ov.get("layout") or {})

    writer = PdfWriter()
    for i, page in enumerate(reader.pages):
        for layout in plan.get(i, []):
            try:
                pw = float(page.mediabox.width)
                ph = float(page.mediabox.height)
                ov_pdf = _render_overlay_pdf_page(layout, ctx, pw, ph)
                ov_page = PdfReader(io.BytesIO(ov_pdf)).pages[0]
                page.merge_page(ov_page)
            except Exception:
                pass
        writer.add_page(page)

    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


# --------------------------------------------------------------------------- #
# PPTX : formes natives via python-pptx
# --------------------------------------------------------------------------- #
def _rgb(color, default="000000"):
    from pptx.dml.color import RGBColor
    try:
        h = str(color or default).lstrip("#")
        if len(h) == 3:
            h = "".join(c * 2 for c in h)
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    except Exception:
        return RGBColor(0, 0, 0)


def _pptx_font_name(key):
    return {
        "title": "Montserrat SemiBold", "heading": "Montserrat SemiBold",
        "subtitle": "Montserrat", "light": "Montserrat Light",
    }.get(key, "Roboto")


def _asset_stream(el, ctx):
    """Renvoie un flux d'image (BytesIO/chemin) pour un élément image/logo."""
    if el.get("type") == "logo":
        p = LR._logo_path(ctx)
        return p or None
    url = el.get("asset_url") or ""
    if url.startswith("data:"):
        try:
            import base64
            return io.BytesIO(base64.b64decode(url.split(",", 1)[1]))
        except Exception:
            return None
    return LR._resolve_asset(url)


def apply_pptx(base_pptx: bytes, document, ctx) -> bytes:
    """Ajoute les éléments des calques sur les diapositives visées d'un PPTX."""
    overlays = _overlays_for(document, "pptx")
    if not overlays:
        return base_pptx
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    prs = Presentation(io.BytesIO(base_pptx))
    slides = list(prs.slides)
    total = len(slides)
    if total == 0:
        return base_pptx

    sw_pt = prs.slide_width / EMU_PER_PT
    sh_pt = prs.slide_height / EMU_PER_PT
    pctx = placeholder_context(ctx)

    def emu(pt):
        return Emu(int(round(pt * EMU_PER_PT)))

    for ov in overlays:
        layout = ov.get("layout") or {}
        idxs = parse_pages((ov.get("target") or {}).get("pages", "all"), total)
        # coordonnées de conception -> points de la diapo réelle
        scaled, fs = _scaled_layout(layout, sw_pt, sh_pt)
        for idx in idxs:
            slide = slides[idx]
            for el in scaled.get("elements", []):
                _add_pptx_element(slide, el, ctx, pctx, emu, Pt, MSO_SHAPE,
                                  MSO_CONNECTOR, PP_ALIGN, MSO_ANCHOR)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _add_pptx_element(slide, el, ctx, pctx, emu, Pt, MSO_SHAPE, MSO_CONNECTOR,
                      PP_ALIGN, MSO_ANCHOR):
    etype = el.get("type", "text")
    x = float(el.get("x", 0)); y = float(el.get("y", 0))
    w = float(el.get("w", 100)); h = float(el.get("h", 20))
    try:
        if etype in ("image", "logo"):
            src = _asset_stream(el, ctx)
            if src is None:
                return
            if el.get("fit", "contain") == "stretch":
                slide.shapes.add_picture(src, emu(x), emu(y), emu(w), emu(h))
            else:
                slide.shapes.add_picture(src, emu(x), emu(y), width=emu(w))

        elif etype == "line":
            conn = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, emu(x), emu(y), emu(x + w), emu(y))
            conn.line.color.rgb = _rgb(el.get("color", "#000000"))
            conn.line.width = Pt(float(el.get("width", 1)))

        elif etype in ("rect", "ellipse"):
            radius = float(el.get("radius", 0) or 0)
            if etype == "ellipse":
                shp_type = MSO_SHAPE.OVAL
            elif radius:
                shp_type = MSO_SHAPE.ROUNDED_RECTANGLE
            else:
                shp_type = MSO_SHAPE.RECTANGLE
            shp = slide.shapes.add_shape(shp_type, emu(x), emu(y), emu(w), emu(h))
            if el.get("fill"):
                shp.fill.solid()
                shp.fill.fore_color.rgb = _rgb(el["fill"])
            else:
                shp.fill.background()
            sw = float(el.get("stroke_width", 0) or 0)
            if el.get("stroke") and sw:
                shp.line.color.rgb = _rgb(el["stroke"])
                shp.line.width = Pt(sw)
            else:
                shp.line.fill.background()
            shp.shadow.inherit = False

        else:  # text
            text = interpolate(el.get("text", ""), pctx)
            tb = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = 0; tf.margin_right = 0
            tf.margin_top = 0; tf.margin_bottom = 0
            align = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(
                el.get("align"), PP_ALIGN.LEFT)
            color = _rgb(el.get("color", "#000000"))
            size = float(el.get("size", 14))
            fname = _pptx_font_name(el.get("font", "body"))
            for i, line in enumerate(str(text).split("\n")):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = align
                run = p.add_run()
                run.text = line
                run.font.size = Pt(size)
                run.font.bold = bool(el.get("bold"))
                run.font.italic = bool(el.get("italic"))
                run.font.name = fname
                run.font.color.rgb = color
    except Exception:
        pass
