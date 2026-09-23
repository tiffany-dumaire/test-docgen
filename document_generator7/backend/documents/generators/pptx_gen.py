"""
Générateur PowerPoint (.pptx) à partir du schéma par blocs.

Un modèle PPTX produit :
  - une diapositive de titre (page de garde) : titre + sous-titre ;
  - une diapositive par bloc « heading » (le titre ouvre une nouvelle diapo),
    les blocs suivants (text, listes, image, table) alimentent son corps.

Styles configurables via ``template.settings["pptx"]`` :
  {theme_color, title_font, body_font, title_size, body_size}
avec repli sur ``settings["table_color"]`` et des valeurs par défaut.
"""
import io

from .base import GenerationContext, interpolate, placeholder_context


def _cfg(ctx):
    s = ctx.document.template.settings or {}
    p = s.get("pptx", {}) or {}
    return {
        "theme_color": (p.get("theme_color") or s.get("table_color") or "4472C4").lstrip("#"),
        "title_font": p.get("title_font", "Calibri Light"),
        "body_font": p.get("body_font", "Calibri"),
        "title_size": int(p.get("title_size", 32)),
        "body_size": int(p.get("body_size", 18)),
        "cover_title": s.get("cover_title", "{{document_title}}"),
        "cover_subtitle": s.get("cover_subtitle", "{{client_name}} — {{project_name}}"),
    }


def _rgb(hexstr):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexstr.lstrip("#").upper())


def render(ctx: GenerationContext) -> bytes:
    from pptx import Presentation
    from pptx.util import Pt, Inches

    pctx = placeholder_context(ctx)
    cfg = _cfg(ctx)
    schema = ctx.document.template.schema or []

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    SW, SH = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]  # disposition vierge
    theme = _rgb(cfg["theme_color"])

    def add_slide():
        return prs.slides.add_slide(blank)

    def band(slide, height=Inches(1.4)):
        from pptx.enum.shapes import MSO_SHAPE
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, height)
        shp.fill.solid(); shp.fill.fore_color.rgb = theme
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def textbox(slide, left, top, width, height):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tb.text_frame.word_wrap = True
        return tb

    def set_run(p, text, size, bold, color, font):
        run = p.add_run(); run.text = text
        run.font.size = Pt(size); run.font.bold = bold
        run.font.name = font
        if color is not None:
            run.font.color.rgb = color

    # --- Diapositive de titre ---
    s0 = add_slide()
    band(s0, Inches(2.2))
    tb = textbox(s0, Inches(0.8), Inches(2.6), Inches(11.7), Inches(2))
    p = tb.text_frame.paragraphs[0]
    set_run(p, interpolate(cfg["cover_title"], pctx), cfg["title_size"] + 12, True, theme, cfg["title_font"])
    p2 = tb.text_frame.add_paragraph()
    set_run(p2, interpolate(cfg["cover_subtitle"], pctx), cfg["body_size"] + 4, False, _rgb("595959"), cfg["body_font"])

    # --- Contenu : une diapo par heading ---
    from pptx.enum.shapes import MSO_SHAPE
    cur = None
    body_tf = None

    def new_content_slide(title):
        nonlocal cur, body_tf
        cur = add_slide()
        bar = cur.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(1.0))
        bar.fill.solid(); bar.fill.fore_color.rgb = theme; bar.line.fill.background()
        bar.shadow.inherit = False
        tt = textbox(cur, Inches(0.6), Inches(0.18), Inches(12), Inches(0.7))
        tp = tt.text_frame.paragraphs[0]
        set_run(tp, title, cfg["title_size"], True, _rgb("FFFFFF"), cfg["title_font"])
        body = textbox(cur, Inches(0.7), Inches(1.3), Inches(12), Inches(5.8))
        body_tf = body.text_frame
        body_tf.word_wrap = True
        return cur

    def ensure_slide():
        nonlocal cur
        if cur is None:
            new_content_slide(interpolate("{{document_title}}", pctx))

    def add_paragraph(text, bullet=False, size=None, bold=False):
        ensure_slide()
        tf = body_tf
        p = tf.paragraphs[0] if (len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs) else tf.add_paragraph()
        prefix = "•  " if bullet else ""
        set_run(p, prefix + text, size or cfg["body_size"], bold, _rgb("1F1F1F"), cfg["body_font"])

    for block in schema:
        bt = block.get("type")
        if bt == "heading":
            new_content_slide(interpolate(block.get("text", ""), pctx))
        elif bt == "text":
            for line in interpolate(block.get("text", ""), pctx).split("\n"):
                if line.strip():
                    add_paragraph(line.strip())
        elif bt in ("bullet_list", "numbered_list"):
            items = block.get("items") or [l for l in interpolate(block.get("text", ""), pctx).split("\n") if l.strip()]
            for i, it in enumerate(items, 1):
                add_paragraph(interpolate(str(it), pctx), bullet=(bt == "bullet_list") or None,
                              size=cfg["body_size"])
        elif bt == "richtext":
            import re
            txt = re.sub(r"<[^>]+>", "", interpolate(block.get("text", ""), pctx))
            for line in txt.split("\n"):
                if line.strip():
                    add_paragraph(line.strip())
        elif bt == "image":
            from .layout_render import _resolve_asset
            path = _resolve_asset(block.get("asset_url"))
            if path and cur is not None:
                try:
                    cur.shapes.add_picture(path, Inches(0.7), Inches(1.4), width=Inches(6))
                except Exception:
                    pass
        elif bt == "table":
            _add_table(cur or new_content_slide(block.get("label", "Tableau")),
                       block, ctx, theme, cfg, Inches, Pt)
        elif bt == "spacer":
            add_paragraph(" ")

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _add_table(slide, block, ctx, theme, cfg, Inches, Pt):
    from pptx.dml.color import RGBColor
    stored = (ctx.data or {}).get(block.get("key", "")) or {}
    columns = stored.get("columns") or [c.get("label", "") for c in block.get("columns", [])]
    rows = stored.get("rows", [])
    if not columns:
        return
    nrows = len(rows) + 1
    ncols = len(columns)
    tbl_shape = slide.shapes.add_table(nrows, ncols, Inches(0.7), Inches(1.5),
                                       Inches(12), Inches(0.4 * nrows))
    table = tbl_shape.table
    for j, col in enumerate(columns):
        cell = table.cell(0, j); cell.text = str(col)
        cell.fill.solid(); cell.fill.fore_color.rgb = theme
        for para in cell.text_frame.paragraphs:
            for r in para.runs:
                r.font.bold = True; r.font.color.rgb = RGBColor.from_string("FFFFFF")
                r.font.size = Pt(12)
    for i, row in enumerate(rows, 1):
        for j in range(ncols):
            val = row[j] if j < len(row) else ""
            c = table.cell(i, j); c.text = "" if val is None else str(val)
            for para in c.text_frame.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(11)
