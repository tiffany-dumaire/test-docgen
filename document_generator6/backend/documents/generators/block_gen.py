"""
Générateur « par blocs », fidèle à Modèle.docx.

- Word : part de Modele.docx (styles, polices Montserrat/Roboto, en-têtes,
  marges) et reconstruit la structure page de garde / suivi / sommaire (champ
  TOC) / contenu.
- PDF : reproduit la même charte (polices embarquées, couleurs, tailles) et la
  même structure, avec sommaire automatique (TableOfContents).

Types de blocs : heading (1-5), text, richtext (HTML), bullet_list,
numbered_list, image, logo, table, field, code, link, contacts, spacer.
"""
import io
import os

from django.conf import settings as dj_settings

from . import housestyle as HS
from .base import GenerationContext, interpolate, placeholder_context
from .htmlparse import inline_runs, parse_html


# ---------------------------------------------------------------------------
# Helpers communs
# ---------------------------------------------------------------------------
def _tpl_settings(ctx):
    s = ctx.document.template.settings or {}
    return {
        "cover_title": s.get("cover_title", "{{document_title}}"),
        "cover_subtitle": s.get("cover_subtitle",
                                "{{client_name}} — {{project_name}}"),
        "include_cover": s.get("include_cover", True),
        "include_suivi": s.get("include_suivi", True),
        "include_toc": s.get("include_toc", True),
        "table_color": (s.get("table_color") or "").lstrip("#") or None,
    }


def _resolve_asset(url):
    """Convertit une URL média en chemin fichier local."""
    if not url:
        return None
    media_url = dj_settings.MEDIA_URL
    idx = url.find(media_url)
    rel = url[idx + len(media_url):] if idx >= 0 else url.lstrip("/")
    path = os.path.join(dj_settings.MEDIA_ROOT, rel)
    return path if os.path.exists(path) else None


def _logo_path(ctx):
    logo = getattr(ctx.company, "logo", None)
    try:
        if logo and logo.path and os.path.exists(logo.path):
            return logo.path
    except Exception:
        pass
    return None


def _field_display(block, ctx):
    key = block.get("key", "")
    value = ctx.data.get(key, "")
    if isinstance(value, (list, dict)):
        value = ""
    value = "" if value is None else str(value)
    if block.get("show_label", True) and block.get("label"):
        return block["label"], value
    return None, value


def _table_data(block, ctx):
    stored = ctx.data.get(block.get("key", "")) or {}
    columns = stored.get("columns") or [c.get("label", "")
                                        for c in block.get("columns", [])]
    rows = stored.get("rows", [])
    width = len(columns)
    norm = [[("" if c is None else str(c)) for c in (list(r) + [""] * width)[:width]]
            for r in rows]
    return columns, norm


# ===========================================================================
# WORD  (base = Modele.docx)
# ===========================================================================
def render_docx(ctx: GenerationContext) -> bytes:
    from docx import Document as Docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Emu, Inches, Pt, RGBColor

    pctx = placeholder_context(ctx)
    cfg = _tpl_settings(ctx)
    table_hex = (cfg.get("table_color") or "1F497D").upper()

    from . import styles as STYLES
    _resolved_styles = STYLES.resolve(ctx)

    def shade_cell(cell, fill_hex, white=True):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), fill_hex)
        tcPr.append(shd)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                if white:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    doc = Docx(HS.BASE_DOCX) if os.path.exists(HS.BASE_DOCX) else Docx()

    # -- Vider le corps en conservant sectPr --
    body = doc.element.body
    sectPr = body.find(qn("w:sectPr"))

    # Largeur utile (EMU) calculée depuis le sectPr avant retrait (1 twip = 635 EMU)
    usable_emu_val = int((11906 - 1134 - 1134) * 635)  # défaut A4 marges modèle
    pg_w_tw, pg_h_tw = 11906, 16838  # A4 par défaut (twips)
    if sectPr is not None:
        pgsz = sectPr.find(qn("w:pgSz"))
        pgmar = sectPr.find(qn("w:pgMar"))
        try:
            w = int(pgsz.get(qn("w:w")))
            pg_w_tw = w
            pg_h_tw = int(pgsz.get(qn("w:h")))
            ml = int(pgmar.get(qn("w:left")))
            mr = int(pgmar.get(qn("w:right")))
            usable_emu_val = int((w - ml - mr) * 635)
        except Exception:
            pass

    # -- Vider le corps SAUF sectPr (nécessaire pour add_table) --
    for child in list(body):
        if child is not sectPr:
            body.remove(child)

    # -- Forcer la mise à jour des champs (sommaire) à l'ouverture --
    try:
        s_el = doc.settings.element
        if s_el.find(qn("w:updateFields")) is None:
            uf = OxmlElement("w:updateFields")
            uf.set(qn("w:val"), "true")
            s_el.append(uf)
    except Exception:
        pass

    styles = {s.style_id: s for s in doc.styles}

    def style(sid):
        return styles.get(sid)

    def align(p, a):
        p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER,
                       "right": WD_ALIGN_PARAGRAPH.RIGHT,
                       "left": WD_ALIGN_PARAGRAPH.LEFT}.get(a, WD_ALIGN_PARAGRAPH.LEFT)

    def para(text="", sid=None):
        p = doc.add_paragraph()
        st = style(sid) if sid else None
        if st is not None:
            p.style = st
        if text:
            p.add_run(text)
        # Surcharges de style en cascade (entreprise → projet → modèle)
        if sid:
            STYLES.apply(p, STYLES.element_for_style_id(sid), _resolved_styles)
        return p

    def page_break():
        from docx.enum.text import WD_BREAK
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def usable_width_emu():
        return usable_emu_val

    def add_hyperlink(paragraph, url, text):
        part = paragraph.part
        r_id = part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True)
        hyper = OxmlElement("w:hyperlink")
        hyper.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        c = OxmlElement("w:color"); c.set(qn("w:val"), "0000FF"); rpr.append(c)
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
        run.append(rpr)
        t = OxmlElement("w:t"); t.text = text; run.append(t)
        hyper.append(run)
        paragraph._p.append(hyper)

    def shaded_code(text):
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "F4F4F4")
        pPr.append(shd)
        run = p.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        return p

    # ---------------- Mises en page libres (garde / suivi) ----------------
    from . import layout_render as LR
    cover_layout = LR.get_layout(ctx, "cover") if LR.has_layout(ctx, "cover") else None
    suivi_layout = LR.get_layout(ctx, "suivi") if LR.has_layout(ctx, "suivi") else None
    _used_layout = bool(cover_layout or suivi_layout)

    def add_fullpage_image(png_bytes):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0); pf.space_after = Pt(0)
        run = p.add_run()
        try:
            run.add_picture(io.BytesIO(png_bytes),
                            width=Emu(int(pg_w_tw * 635)),
                            height=Emu(int(pg_h_tw * 635)))
        except Exception:
            pass

    def append_image_section_break():
        """Clôt la section « images » : nouvelle section 0 marge, sans en-tête."""
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        sp = OxmlElement("w:sectPr")
        pgSz = OxmlElement("w:pgSz")
        pgSz.set(qn("w:w"), str(pg_w_tw)); pgSz.set(qn("w:h"), str(pg_h_tw))
        sp.append(pgSz)
        pgMar = OxmlElement("w:pgMar")
        for k in ("top", "right", "bottom", "left", "header", "footer", "gutter"):
            pgMar.set(qn(f"w:{k}"), "0")
        sp.append(pgMar)
        # pas de headerReference/footerReference -> aucune en-tête sur ces pages
        titlePg = OxmlElement("w:titlePg"); sp.append(titlePg)
        pPr.append(sp)

    # ---------------- Page de garde ----------------
    if cfg["include_cover"]:
        if cover_layout:
            png, _w, _h = LR.render_png(cover_layout, ctx)
            if png:
                add_fullpage_image(png)
            else:
                para(interpolate(cfg["cover_title"], pctx), HS.DOCX_STYLE_IDS["title"])
        else:
            para(interpolate(cfg["cover_title"], pctx), HS.DOCX_STYLE_IDS["title"])
            para(interpolate(cfg["cover_subtitle"], pctx), HS.DOCX_STYLE_IDS["subtitle"])

    # ---------------- Page de suivi ----------------
    if cfg["include_suivi"] and suivi_layout:
        if cover_layout:
            page_break()
        png, _w, _h = LR.render_png(suivi_layout, ctx)
        if png:
            add_fullpage_image(png)
        if _used_layout:
            append_image_section_break()
    elif cfg["include_suivi"]:
        if _used_layout:
            append_image_section_break()
        page_break()
        para("Identification", HS.DOCX_STYLE_IDS["section"])
        ident = [
            ("Client", ctx.project.client_name),
            ("Projet", ctx.project.name),
            ("Séquence / Référence", ctx.project.reference or "—"),
            ("Confidentialité", ctx.confidentiality_label),
        ]
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for k, v in ident:
            cells = t.add_row().cells
            cells[0].text = k
            cells[1].text = str(v)
            if cells[0].paragraphs[0].runs:
                cells[0].paragraphs[0].runs[0].bold = True

        para("Révisions", HS.DOCX_STYLE_IDS["section"])
        from documents.models import ConfidentialityLevel
        conf_map = dict(ConfidentialityLevel.choices)
        rt = doc.add_table(rows=1, cols=4)
        rt.style = "Table Grid"
        for i, h in enumerate(["Version", "Date", "Par", "Commentaire"]):
            rt.rows[0].cells[i].text = h
            shade_cell(rt.rows[0].cells[i], table_hex)
        versions = list(ctx.document.versions.order_by("version_number"))
        if not versions:
            cells = rt.add_row().cells
            cells[0].text = f"{ctx.version_number}.0"
            cells[1].text = "en cours"
            cells[2].text = ctx.author_initials
            cells[3].text = ctx.comment or "Génération initiale."
        for v in versions:
            cells = rt.add_row().cells
            cells[0].text = f"{v.version_number}.0"
            cells[1].text = v.created_at.strftime("%Y-%m-%d")
            cells[2].text = v.author_initials
            cells[3].text = v.comment or "—"

    # Sécurité : si une garde personnalisée existe mais pas de page de suivi,
    # clôturer ici la section « images » (0 marge, sans en-tête).
    if _used_layout and not cfg["include_suivi"]:
        append_image_section_break()

    # ---------------- Sommaire ----------------
    if cfg["include_toc"]:
        page_break()
        para("Table des matières", HS.DOCX_STYLE_IDS["section"])
        p = doc.add_paragraph()
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), r' TOC \o "1-5" \h \z \u ')
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "Faites clic droit puis « Mettre à jour les champs » (ou F9)."
        r.append(t); fld.append(r)
        p._p.append(fld)
        page_break()

    # ---------------- Contenu ----------------
    def render_richtext(html):
        for node in parse_html(html):
            if node["kind"] == "heading":
                p = para("", HS.DOCX_STYLE_IDS.get(f"h{node['level']}", "Titre2"))
                _add_inline(p, node["html"], doc, add_hyperlink)
            elif node["kind"] == "paragraph":
                p = para("", HS.DOCX_STYLE_IDS["normal"])
                _add_inline(p, node["html"], doc, add_hyperlink)
            elif node["kind"] == "list":
                for item in node["items"]:
                    sid = _list_style(styles, node["ordered"])
                    p = doc.add_paragraph()
                    if sid:
                        p.style = styles[sid]
                    else:
                        p.add_run("• " if not node["ordered"] else "– ")
                    _add_inline(p, item, doc, add_hyperlink)
            elif node["kind"] == "code":
                shaded_code(node["text"])

    for block in ctx.document.template.schema:
        bt = block.get("type")
        if bt == "heading":
            lvl = int(block.get("level", 2))
            p = para(interpolate(block.get("text", ""), pctx),
                     HS.DOCX_STYLE_IDS.get(f"h{lvl}", "Titre2"))
        elif bt == "text":
            para(interpolate(block.get("text", ""), pctx),
                 HS.DOCX_STYLE_IDS["normal"])
        elif bt == "richtext":
            render_richtext(interpolate(block.get("text", ""), pctx))
        elif bt in ("bullet_list", "numbered_list"):
            ordered = bt == "numbered_list"
            items = block.get("items")
            if not items:
                items = [l for l in interpolate(block.get("text", ""), pctx).split("\n") if l.strip()]
            for n, it in enumerate(items, 1):
                sid = _list_style(styles, ordered)
                p = doc.add_paragraph()
                if sid:
                    p.style = styles[sid]
                    p.add_run(interpolate(str(it), pctx))
                else:
                    p.add_run((f"{n}. " if ordered else "• ") + interpolate(str(it), pctx))
        elif bt == "diagram":
            from . import diagram_render as DR
            png, wpx, hpx = DR.render_png(block, ctx, color=table_hex)
            if png and wpx:
                pct = float(block.get("width_pct", 100)) / 100.0
                width_emu = int(usable_width_emu() * max(0.3, min(pct, 1.0)))
                height_emu = int(width_emu * hpx / wpx)
                p = doc.add_paragraph()
                align(p, block.get("align", "center"))
                run = p.add_run()
                try:
                    run.add_picture(io.BytesIO(png), width=Emu(width_emu), height=Emu(height_emu))
                except Exception:
                    p.add_run("[diagramme]")
        elif bt in ("image", "logo"):
            path = _logo_path(ctx) if bt == "logo" else _resolve_asset(block.get("asset_url"))
            if path:
                pct = float(block.get("width_pct", 40)) / 100.0
                width = Emu(int(usable_width_emu() * max(0.05, min(pct, 1.0))))
                p = doc.add_paragraph()
                align(p, block.get("align", "left"))
                run = p.add_run()
                try:
                    run.add_picture(path, width=width)
                except Exception:
                    p.add_run("[image]")
        elif bt == "table":
            columns, rows = _table_data(block, ctx)
            if block.get("label"):
                para(block["label"], HS.DOCX_STYLE_IDS["h3"])
            if columns:
                tt = doc.add_table(rows=1, cols=len(columns))
                tt.style = "Table Grid"
                for i, col in enumerate(columns):
                    tt.rows[0].cells[i].text = str(col)
                    shade_cell(tt.rows[0].cells[i], table_hex)
                for r in rows:
                    cells = tt.add_row().cells
                    for i, val in enumerate(r):
                        cells[i].text = str(val)
        elif bt == "field":
            label, value = _field_display(block, ctx)
            p = doc.add_paragraph()
            p.style = styles.get(HS.DOCX_STYLE_IDS["normal"])
            if label:
                p.add_run(f"{label} : ").bold = True
            p.add_run(value)
        elif bt == "code":
            shaded_code(interpolate(block.get("text", ""), pctx))
        elif bt == "link":
            url = interpolate(block.get("url", ""), pctx)
            label = interpolate(block.get("label", "") or url, pctx)
            p = doc.add_paragraph()
            if url:
                add_hyperlink(p, url, label)
            else:
                p.add_run(label)
        elif bt == "contacts":
            _contacts_docx(doc, block, ctx, styles)
        elif bt == "spacer":
            doc.add_paragraph()

    # -- Pied : profil entreprise --
    foot = doc.add_paragraph()
    align(foot, "center")
    bits = [b for b in [ctx.company.website_url, ctx.company.email,
                        ctx.company.phone] if b]
    if bits:
        fr = foot.add_run("  ·  ".join(bits))
        fr.font.size = Pt(8)
        fr.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    if sectPr is not None:
        body.append(sectPr)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _list_style(styles, ordered):
    wanted = "ListNumber" if ordered else "ListBullet"
    if wanted in styles:
        return wanted
    for sid, st in styles.items():
        nm = (st.name or "").lower()
        if ordered and "number" in nm and "list" in nm:
            return sid
        if not ordered and "bullet" in nm and "list" in nm:
            return sid
    return None


def _add_inline(paragraph, html, doc, add_hyperlink):
    from docx.shared import RGBColor
    for seg in inline_runs(html):
        if seg.get("href"):
            add_hyperlink(paragraph, seg["href"], seg["text"])
            continue
        run = paragraph.add_run(seg["text"])
        run.bold = seg.get("bold", False)
        run.italic = seg.get("italic", False)
        run.underline = seg.get("underline", False)


def _contacts_docx(doc, block, ctx, styles):
    scope = block.get("scope", "both")
    contacts = []
    if scope in ("client", "both"):
        contacts += [("Client", c) for c in ctx.client_contacts()]
    if scope in ("internal", "both"):
        contacts += [("Interne", c) for c in ctx.internal_contacts()]
    p = doc.add_paragraph()
    p.style = styles.get(HS.DOCX_STYLE_IDS["h3"])
    p.add_run(block.get("label", "Contacts"))
    if not contacts:
        doc.add_paragraph("Aucun contact renseigné.")
        return
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for i, h in enumerate(["Rôle", "Nom", "Fonction", "Email", "Téléphone"]):
        t.rows[0].cells[i].text = h
        if t.rows[0].cells[i].paragraphs[0].runs:
            t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for kind, c in contacts:
        cells = t.add_row().cells
        for i, v in enumerate([kind, c.full_name, c.role or "—",
                               c.email or "—", c.phone or "—"]):
            cells[i].text = v


# ===========================================================================
# PDF  (mêmes polices / couleurs / structure que le modèle)
# ===========================================================================
def _hex(c):
    from reportlab.lib.colors import HexColor
    return HexColor(c)


def _sanitize_inline(html):
    """Adapte le HTML en ligne au mini-langage de reportlab."""
    return (html.replace("<strong>", "<b>").replace("</strong>", "</b>")
            .replace("<em>", "<i>").replace("</em>", "</i>"))


def render_pdf(ctx: GenerationContext) -> bytes:
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm, mm
    from reportlab.platypus import (BaseDocTemplate, Frame, Image, ListFlowable,
                                    ListItem, NextPageTemplate, PageBreak,
                                    PageTemplate, Paragraph, Spacer, Table,
                                    TableStyle)
    from reportlab.platypus.tableofcontents import TableOfContents
    from reportlab.lib.utils import ImageReader

    HS.register_pdf_fonts()
    pctx = placeholder_context(ctx)
    cfg = _tpl_settings(ctx)
    table_color = _hex('#' + cfg['table_color']) if cfg.get('table_color') else _hex(HS.COLORS['primary'])
    W, H = A4

    body_font = "Roboto" if HS.font_available("Roboto") else "Helvetica"
    ms_sb = "Montserrat-SemiBold" if HS.font_available("Montserrat-SemiBold") else "Helvetica-Bold"
    ms_rg = "Montserrat" if HS.font_available("Montserrat") else "Helvetica"
    ms_lt = "Montserrat-Light" if HS.font_available("Montserrat-Light") else "Helvetica"

    def pstyle(name, font, size, color, **kw):
        kw.setdefault("leading", size * 1.3)
        return ParagraphStyle(name, fontName=font, fontSize=size,
                              textColor=_hex(color), **kw)

    body_style = pstyle("DGbody", body_font, HS.BODY[1], HS.BODY[2], spaceAfter=5)
    title_style = pstyle("DGtitle", ms_sb, HS.TITLE[1], HS.TITLE[2], leading=32)
    sub_style = pstyle("DGsub", ms_rg, HS.SUBTITLE[1], HS.SUBTITLE[2], leading=22)
    section_style = pstyle("DGsection", ms_sb, HS.SECTION[1], HS.SECTION[2],
                           spaceBefore=6, spaceAfter=10)
    heading_styles = {}
    for lvl, (f, sz, col) in HS.HEADINGS.items():
        fnt = {1: ms_sb, 2: ms_sb, 3: ms_rg, 4: ms_lt, 5: ms_rg}[lvl]
        heading_styles[lvl] = pstyle(f"DGH{lvl}", fnt, sz, col,
                                     spaceBefore=12 - lvl, spaceAfter=5)
    code_style = pstyle("DGcode", "Courier", 8.5, "#1E293B",
                        backColor=_hex(HS.COLORS["code_bg"]), borderPadding=6,
                        leading=12)
    link_style = pstyle("DGlink", body_font, HS.BODY[1], HS.COLORS["link"])

    # ---- Modèle de page ----
    class Doc(BaseDocTemplate):
        def afterFlowable(self, flowable):
            if isinstance(flowable, Paragraph):
                st = flowable.style.name
                if st in ("DGH1", "DGH2", "DGH3", "DGH4", "DGH5"):
                    level = int(st[-1]) - 1
                    self.notify("TOCEntry", (level, flowable.getPlainText(),
                                             self.page))

    def cover_page(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(_hex(HS.COLORS["primary"]))
        canvas.setLineWidth(2)
        canvas.line(2 * cm, H - 4 * cm, 8 * cm, H - 4 * cm)
        canvas.restoreState()

    def content_page(canvas, doc):
        canvas.saveState()
        # bandeau haut discret
        canvas.setFillColor(_hex(HS.COLORS["primary"]))
        canvas.rect(0, H - 0.9 * cm, W, 0.9 * cm, fill=1, stroke=0)
        canvas.setFillColor(_hex("#FFFFFF"))
        canvas.setFont(ms_sb, 9)
        canvas.drawString(2 * cm, H - 0.62 * cm, ctx.company.name or "")
        canvas.setFont(body_font, 7)
        canvas.drawRightString(W - 2 * cm, H - 0.6 * cm,
                               ctx.confidentiality_label.upper())
        # pied
        canvas.setStrokeColor(_hex(HS.COLORS["border"]))
        canvas.line(2 * cm, 1.4 * cm, W - 2 * cm, 1.4 * cm)
        canvas.setFillColor(_hex(HS.COLORS["muted"]))
        canvas.setFont(body_font, 7)
        foot = "  ·  ".join([b for b in [ctx.company.website_url,
                            ctx.company.email] if b])
        canvas.drawString(2 * cm, 1.0 * cm, foot)
        canvas.drawRightString(W - 2 * cm, 1.0 * cm,
                               f"v{ctx.version_number}  ·  Page {doc.page}")
        # filigrane
        if ctx.confidentiality in ("confidential", "restricted"):
            canvas.saveState()
            canvas.translate(W / 2, H / 2); canvas.rotate(45)
            canvas.setFont(ms_sb, 58)
            canvas.setFillColorRGB(0.92, 0.92, 0.92)
            canvas.drawCentredString(0, 0, ctx.confidentiality_label.upper())
            canvas.restoreState()
        canvas.restoreState()

    buffer = io.BytesIO()
    doc = Doc(buffer, pagesize=A4, topMargin=2.4 * cm, bottomMargin=2 * cm,
              leftMargin=2 * cm, rightMargin=2 * cm)
    from . import layout_render as LR
    cover_layout = LR.get_layout(ctx, "cover") if LR.has_layout(ctx, "cover") else None
    suivi_layout = LR.get_layout(ctx, "suivi") if LR.has_layout(ctx, "suivi") else None

    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="f")
    fp_frame = Frame(0, 0, W, H, leftPadding=0, rightPadding=0, topPadding=0,
                     bottomPadding=0, id="fp")
    _templates = []
    if cover_layout and cfg["include_cover"]:
        _templates.append(PageTemplate(id="cover", frames=[fp_frame]))
    else:
        _templates.append(PageTemplate(id="cover", frames=[frame], onPage=cover_page))
    _templates.append(PageTemplate(id="content", frames=[frame], onPage=content_page))
    _templates.append(PageTemplate(id="fullpage", frames=[fp_frame]))
    doc.addPageTemplates(_templates)

    def _full_bleed(layout):
        png, _wpx, _hpx = LR.render_png(layout, ctx)
        if not png:
            return None
        return Image(io.BytesIO(png), width=W, height=H)

    story = []
    usable = doc.width

    # ---- Page de garde ----
    if cfg["include_cover"]:
        if cover_layout:
            _img = _full_bleed(cover_layout)
            if _img is not None:
                story.append(_img)
            else:
                story.append(Spacer(1, 6 * cm))
                story.append(Paragraph(interpolate(cfg["cover_title"], pctx), title_style))
        else:
            story.append(Spacer(1, 6 * cm))
            story.append(Paragraph(interpolate(cfg["cover_title"], pctx), title_style))
            story.append(Spacer(1, 6))
            story.append(Paragraph(interpolate(cfg["cover_subtitle"], pctx), sub_style))

    # ---- Page de suivi ----
    def kv_table(rows):
        data = [[Paragraph(f"<b>{k}</b>", body_style), Paragraph(str(v), body_style)]
                for k, v in rows]
        t = Table(data, colWidths=[5 * cm, usable - 5 * cm])
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, _hex(HS.COLORS["border"])),
            ("BACKGROUND", (0, 0), (0, -1), _hex(HS.COLORS["light"])),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ]))
        return t

    if cfg["include_suivi"] and suivi_layout:
        story.append(NextPageTemplate("fullpage"))
        story.append(PageBreak())
        _img = _full_bleed(suivi_layout)
        if _img is not None:
            story.append(_img)
    elif cfg["include_suivi"]:
        story.append(NextPageTemplate("content"))
        story.append(PageBreak())
        story.append(Paragraph("Identification", section_style))
        story.append(kv_table([
            ("Client", ctx.project.client_name),
            ("Projet", ctx.project.name),
            ("Séquence / Référence", ctx.project.reference or "—"),
            ("Confidentialité", ctx.confidentiality_label),
        ]))
        story.append(Spacer(1, 14))
        story.append(Paragraph("Révisions", section_style))
        header = [["Version", "Date", "Par", "Commentaire"]]
        rows = []
        versions = list(ctx.document.versions.order_by("version_number"))
        if not versions:
            rows.append([f"{ctx.version_number}.0", "en cours", ctx.author_initials,
                         Paragraph(ctx.comment or "Génération initiale.", body_style)])
        for v in versions:
            rows.append([f"{v.version_number}.0", v.created_at.strftime("%Y-%m-%d"),
                         v.author_initials, Paragraph(v.comment or "—", body_style)])
        t = Table(header + rows, colWidths=[2 * cm, 2.6 * cm, 1.6 * cm, usable - 6.2 * cm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), table_color),
            ("TEXTCOLOR", (0, 0), (-1, 0), _hex("#FFFFFF")),
            ("FONTNAME", (0, 0), (-1, 0), ms_sb),
            ("FONTNAME", (0, 1), (-1, -1), body_font),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.4, _hex(HS.COLORS["border"])),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [_hex("#FFFFFF"), _hex(HS.COLORS["light"])]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(t)

    # ---- Sommaire ----
    if cfg["include_toc"]:
        story.append(NextPageTemplate("content"))
        story.append(PageBreak())
        story.append(Paragraph("Table des matières", section_style))
        toc = TableOfContents()
        toc.levelStyles = [
            pstyle("TOC1", ms_sb, 11, HS.COLORS["toc_heading"], leftIndent=0,
                   firstLineIndent=0, spaceBefore=4),
            pstyle("TOC2", ms_rg, 10, "#000000", leftIndent=14, spaceBefore=2),
            pstyle("TOC3", ms_rg, 10, "#000000", leftIndent=28),
            pstyle("TOC4", ms_lt, 10, "#000000", leftIndent=42),
            pstyle("TOC5", ms_rg, 10, HS.COLORS["heading_deep"], leftIndent=56),
        ]
        story.append(toc)

    if cfg["include_cover"] or cfg["include_suivi"] or cfg["include_toc"]:
        story.append(NextPageTemplate("content"))
        story.append(PageBreak())

    # ---- Contenu ----
    def table_flowable(columns, rows):
        header = [[Paragraph(f"<b>{c}</b>", ParagraphStyle('th', parent=body_style, textColor=_hex('#FFFFFF')))
                   for c in columns]]
        data = header + [[Paragraph(str(c), body_style) for c in r] for r in rows]
        t = Table(data, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), table_color),
            ("FONTNAME", (0, 0), (-1, -1), body_font),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.4, _hex(HS.COLORS["border"])),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [_hex("#FFFFFF"), _hex(HS.COLORS["light"])]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ]))
        return t

    def image_flowable(path, block):
        try:
            iw, ih = ImageReader(path).getSize()
            pct = float(block.get("width_pct", 40)) / 100.0
            w = usable * max(0.05, min(pct, 1.0))
            h = w * ih / iw
            img = Image(path, width=w, height=h)
            img.hAlign = block.get("align", "left").upper()
            return img
        except Exception:
            return Paragraph("[image]", body_style)

    def list_flowable(items, ordered):
        flow = [ListItem(Paragraph(_sanitize_inline(str(it)), body_style))
                for it in items]
        return ListFlowable(flow, bulletType="1" if ordered else "bullet",
                            start="1" if ordered else "•", leftIndent=16)

    for block in ctx.document.template.schema:
        bt = block.get("type")
        if bt == "heading":
            lvl = int(block.get("level", 2))
            story.append(Paragraph(interpolate(block.get("text", ""), pctx),
                                   heading_styles.get(lvl, heading_styles[2])))
        elif bt == "text":
            txt = interpolate(block.get("text", ""), pctx).replace("\n", "<br/>")
            story.append(Paragraph(txt, body_style))
        elif bt == "richtext":
            for node in parse_html(interpolate(block.get("text", ""), pctx)):
                if node["kind"] == "heading":
                    story.append(Paragraph(_sanitize_inline(node["html"]),
                                           heading_styles.get(node["level"], heading_styles[2])))
                elif node["kind"] == "paragraph":
                    story.append(Paragraph(_sanitize_inline(node["html"]), body_style))
                elif node["kind"] == "list":
                    story.append(list_flowable([_sanitize_inline(i) for i in node["items"]],
                                               node["ordered"]))
                elif node["kind"] == "code":
                    story.append(Paragraph(node["text"].replace("\n", "<br/>"), code_style))
        elif bt in ("bullet_list", "numbered_list"):
            items = block.get("items") or [l for l in interpolate(block.get("text", ""), pctx).split("\n") if l.strip()]
            items = [interpolate(str(i), pctx) for i in items]
            if items:
                story.append(list_flowable(items, bt == "numbered_list"))
        elif bt == "diagram":
            from . import diagram_render as DR
            png, wpx, hpx = DR.render_png(block, ctx, color=(cfg.get("table_color") or "1F497D"))
            if png and wpx:
                pct = float(block.get("width_pct", 100)) / 100.0
                w = usable * max(0.3, min(pct, 1.0))
                story.append(Image(io.BytesIO(png), width=w, height=w * hpx / wpx))
        elif bt in ("image", "logo"):
            path = _logo_path(ctx) if bt == "logo" else _resolve_asset(block.get("asset_url"))
            if path:
                story.append(image_flowable(path, block))
        elif bt == "table":
            columns, rows = _table_data(block, ctx)
            if block.get("label"):
                story.append(Paragraph(block["label"], heading_styles[3]))
            if columns:
                story.append(table_flowable(columns, rows or [[""] * len(columns)]))
        elif bt == "field":
            label, value = _field_display(block, ctx)
            value = value.replace("\n", "<br/>")
            story.append(Paragraph(f"<b>{label} :</b> {value}" if label else value, body_style))
        elif bt == "code":
            story.append(Paragraph(interpolate(block.get("text", ""), pctx).replace("\n", "<br/>"), code_style))
        elif bt == "link":
            url = interpolate(block.get("url", ""), pctx)
            label = interpolate(block.get("label", "") or url, pctx)
            story.append(Paragraph(f'<a href="{url}"><u>{label}</u></a>', link_style))
        elif bt == "contacts":
            story += _contacts_pdf(block, ctx, heading_styles[3], body_font, ms_sb, table_color)
        elif bt == "spacer":
            story.append(Spacer(1, 12))

    doc.multiBuild(story)
    return buffer.getvalue()


def _contacts_pdf(block, ctx, heading_style, body_font, bold_font, table_color=None):
    table_color = table_color or _hex(HS.COLORS["primary"])
    from reportlab.platypus import Paragraph, Table, TableStyle
    from reportlab.lib.units import cm
    scope = block.get("scope", "both")
    contacts = []
    if scope in ("client", "both"):
        contacts += [("Client", c) for c in ctx.client_contacts()]
    if scope in ("internal", "both"):
        contacts += [("Interne", c) for c in ctx.internal_contacts()]
    out = [Paragraph(block.get("label", "Contacts"), heading_style)]
    if not contacts:
        return out
    header = [["Rôle", "Nom", "Fonction", "Email", "Téléphone"]]
    rows = [[k, c.full_name, c.role or "—", c.email or "—", c.phone or "—"]
            for k, c in contacts]
    t = Table(header + rows, colWidths=[2 * cm, 3.4 * cm, 3 * cm, 4.2 * cm, 2.9 * cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), table_color),
        ("TEXTCOLOR", (0, 0), (-1, 0), _hex("#FFFFFF")),
        ("FONTNAME", (0, 0), (-1, 0), bold_font),
        ("FONTNAME", (0, 1), (-1, -1), body_font),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, _hex(HS.COLORS["border"])),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [_hex("#FFFFFF"), _hex(HS.COLORS["light"])]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    out.append(t)
    return out


# ===========================================================================
# EXCEL
# ===========================================================================
def render_xlsx(ctx: GenerationContext) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HEADER_FILL = PatternFill("solid", fgColor="1F497D")
    ALT_FILL = PatternFill("solid", fgColor="F2F2F2")
    HEADER_FONT = Font(bold=True, color="FFFFFF")
    LABEL_FONT = Font(bold=True, color="1F497D")
    TITLE_FONT = Font(size=16, bold=True, color="1F497D")
    THIN = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    pctx = placeholder_context(ctx)

    wb = Workbook()
    ws = wb.active
    ws.title = "Document"
    row = 1

    def autosize(sheet):
        for col in sheet.columns:
            length, letter = 0, None
            for c in col:
                if letter is None:
                    letter = get_column_letter(c.column)
                if c.value is not None:
                    length = max(length, len(str(c.value)))
            if letter:
                sheet.column_dimensions[letter].width = min(length + 4, 70)

    for block in ctx.document.template.schema:
        bt = block.get("type")
        if bt in ("heading", "text", "richtext", "link", "code"):
            txt = interpolate(block.get("text", "") or block.get("label", ""), pctx)
            if bt == "link":
                txt = interpolate(block.get("label", "") or block.get("url", ""), pctx)
            cell = ws.cell(row=row, column=1, value=txt)
            cell.font = TITLE_FONT if bt == "heading" and int(block.get("level", 2)) == 1 else LABEL_FONT if bt == "heading" else Font()
            row += 2 if bt == "heading" else 1
        elif bt in ("bullet_list", "numbered_list"):
            items = block.get("items") or [l for l in interpolate(block.get("text", ""), pctx).split("\n") if l.strip()]
            for n, it in enumerate(items, 1):
                ws.cell(row=row, column=1,
                        value=(f"{n}. " if bt == "numbered_list" else "• ") + interpolate(str(it), pctx))
                row += 1
            row += 1
        elif bt == "field":
            label, value = _field_display(block, ctx)
            if label:
                ws.cell(row=row, column=1, value=label).font = LABEL_FONT
                ws.cell(row=row, column=2, value=value)
            else:
                ws.cell(row=row, column=1, value=value)
            row += 1
        elif bt == "table":
            columns, rows = _table_data(block, ctx)
            if block.get("label"):
                ws.cell(row=row, column=1, value=block["label"]).font = LABEL_FONT
                row += 1
            for j, col in enumerate(columns, 1):
                c = ws.cell(row=row, column=j, value=str(col))
                c.fill = HEADER_FILL; c.font = HEADER_FONT; c.border = BORDER
                c.alignment = Alignment(horizontal="center")
            row += 1
            for r in rows:
                for j, val in enumerate(r, 1):
                    c = ws.cell(row=row, column=j, value=val)
                    c.border = BORDER
                    if row % 2 == 0:
                        c.fill = ALT_FILL
                row += 1
            row += 1
        elif bt == "contacts":
            scope = block.get("scope", "both")
            cc = []
            if scope in ("client", "both"):
                cc += [("Client", c) for c in ctx.client_contacts()]
            if scope in ("internal", "both"):
                cc += [("Interne", c) for c in ctx.internal_contacts()]
            ws.cell(row=row, column=1, value=block.get("label", "Contacts")).font = LABEL_FONT
            row += 1
            for j, h in enumerate(["Rôle", "Nom", "Fonction", "Email", "Téléphone"], 1):
                c = ws.cell(row=row, column=j, value=h)
                c.fill = HEADER_FILL; c.font = HEADER_FONT; c.border = BORDER
            row += 1
            for kind, c in cc:
                for j, v in enumerate([kind, c.full_name, c.role, c.email, c.phone], 1):
                    ws.cell(row=row, column=j, value=v).border = BORDER
                row += 1
            row += 1
        elif bt == "spacer":
            row += 1

    autosize(ws)

    # Feuille Historique
    hist = wb.create_sheet("Historique")
    from documents.models import ConfidentialityLevel
    conf_map = dict(ConfidentialityLevel.choices)
    for j, h in enumerate(["Version", "Date", "Auteur", "Confidentialité", "Commentaire"], 1):
        c = hist.cell(row=1, column=j, value=h)
        c.fill = HEADER_FILL; c.font = HEADER_FONT
    r = 2
    for v in ctx.document.versions.order_by("version_number"):
        hist.append([f"v{v.version_number}", v.created_at.strftime("%d/%m/%Y %H:%M"),
                     v.author_initials, conf_map.get(v.confidentiality, v.confidentiality),
                     v.comment])
        r += 1
    if r == 2:
        hist.append([f"v{ctx.version_number}", "en cours", ctx.author_initials,
                     ctx.confidentiality_label, ctx.comment])
    autosize(hist)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


BLOCK_RENDERERS = {
    "pdf": render_pdf,
    "docx": render_docx,
    "xlsx": render_xlsx,
}
