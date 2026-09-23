"""
Générateur de classeurs Excel personnalisables.

Un modèle Excel « classeur » décrit plusieurs onglets dans
`template.settings["excel"]["sheets"]`. Trois types d'onglets :

- ``table``  : tableau de données à colonnes typées, avec regroupements
               d'en-têtes sur deux lignes (ligne 1 = groupes fusionnés,
               ligne 2 = colonnes, ligne 3+ = données).
- ``pivot``  : tableau croisé calculé à partir d'un onglet ``table``.
- ``info``   : informations fixes typées (date, texte, monnaie,
               nombre entier / à virgule, pourcentage).

Les données saisies pour le document se trouvent dans
``document.data["sheets"][<id_onglet>]`` :
``{"columns": [...], "rows": [[...], ...]}``.
"""
from datetime import date, datetime

from .base import GenerationContext, interpolate, placeholder_context

# Couleurs de la charte
PRIMARY = "1F497D"
GROUP_FILL = "1F497D"
HEADER_FILL = "4472C4"
ALT_FILL = "EDF2FA"
LIGHT = "F5F7FB"
BORDER_COLOR = "BFC7D2"

NUMBER_FORMATS = {
    "text": None,
    "integer": "#,##0",
    "decimal": "#,##0.00",
    "percent": "0.00%",
    "date": "dd/mm/yyyy",
}


def _fmt(col):
    """Format Excel d'après le type de colonne / champ."""
    t = col.get("type", "text")
    if t == "currency":
        symbol = col.get("symbol", "€")
        dec = int(col.get("decimals", 2))
        zeros = "0" * dec
        return f'#,##0.{zeros} "{symbol}"' if dec else f'#,##0 "{symbol}"'
    if t == "decimal":
        dec = int(col.get("decimals", 2))
        return "#,##0." + ("0" * dec) if dec else "#,##0"
    return NUMBER_FORMATS.get(t)


def _coerce(value, col, pctx):
    """Convertit une valeur (souvent une chaîne) selon le type attendu."""
    t = col.get("type", "text")
    if isinstance(value, str):
        value = interpolate(value, pctx)
    if value in (None, ""):
        return None
    if t in ("integer", "decimal", "currency", "percent"):
        if isinstance(value, (int, float)):
            num = float(value)
        else:
            cleaned = (str(value).replace(" ", "").replace("\u202f", "")
                       .replace("€", "").replace("%", "").replace(",", "."))
            try:
                num = float(cleaned)
            except ValueError:
                return str(value)
        if t == "integer":
            return int(round(num))
        if t == "percent":
            # une valeur "20" ou "0.2" -> 20 % / 20 %
            return num / 100.0 if abs(num) > 1 else num
        return num
    if t == "date":
        if isinstance(value, (date, datetime)):
            return value
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(str(value), fmt).date()
            except ValueError:
                continue
        return str(value)
    return str(value)


def _sheet_name(name, used):
    """Nom d'onglet valide (<=31 car., caractères interdits retirés, unique)."""
    for ch in "[]:*?/\\":
        name = name.replace(ch, " ")
    name = (name.strip() or "Feuille")[:31]
    base, i = name, 2
    while name.lower() in used:
        suffix = f" {i}"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name.lower())
    return name


def _effective_columns(sdef, sheet_data):
    """Colonnes réellement utilisées (modèle + colonnes ajoutées au remplissage)."""
    tpl_cols = list(sdef.get("columns", []))
    data_cols = (sheet_data or {}).get("columns")
    if not data_cols:
        return tpl_cols
    # data_cols peut être une liste de libellés ou de dicts
    merged = []
    for i, dc in enumerate(data_cols):
        if isinstance(dc, dict):
            base = dict(tpl_cols[i]) if i < len(tpl_cols) else {}
            base.update(dc)
            merged.append(base)
        else:  # libellé simple
            if i < len(tpl_cols):
                c = dict(tpl_cols[i]); c["label"] = dc; merged.append(c)
            else:
                merged.append({"label": str(dc), "type": "text",
                               "key": f"col{i + 1}"})
    return merged


def render(ctx: GenerationContext) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    pctx = placeholder_context(ctx)
    excel_cfg = (ctx.document.template.settings or {}).get("excel", {})
    sheets = excel_cfg.get("sheets", [])
    data_sheets = (ctx.data or {}).get("sheets", {})

    thin = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    group_font = Font(bold=True, color="FFFFFF", size=11)
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, color=PRIMARY, size=14)
    label_font = Font(bold=True, color=PRIMARY)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")

    def autosize(ws, ncols, minw=10, maxw=48):
        for c in range(1, ncols + 1):
            letter = get_column_letter(c)
            longest = minw
            for cell in ws[letter]:
                if cell.value is not None:
                    longest = max(longest, len(str(cell.value)) + 2)
            ws.column_dimensions[letter].width = min(longest, maxw)

    wb = Workbook()
    wb.remove(wb.active)
    used_names = set()

    # 1) Pré-calcule les données brutes des onglets « table » (pour les pivots)
    tables = {}
    for sdef in sheets:
        if sdef.get("type") == "table":
            cols = _effective_columns(sdef, data_sheets.get(sdef.get("id"), {}))
            raw_rows = (data_sheets.get(sdef.get("id"), {}) or {}).get("rows", [])
            coerced = []
            for r in raw_rows:
                row = list(r) + [None] * (len(cols) - len(r))
                coerced.append([_coerce(row[i], cols[i], pctx)
                                for i in range(len(cols))])
            tables[sdef.get("id")] = {"def": sdef, "columns": cols, "rows": coerced}

    # 2) Construit les onglets dans l'ordre défini
    for sdef in sheets:
        stype = sdef.get("type", "table")
        ws = wb.create_sheet(_sheet_name(sdef.get("name", "Feuille"), used_names))

        # Titre d'onglet (facultatif)
        title = interpolate(sdef.get("title", ""), pctx)
        start_row = 1
        if title:
            ws.cell(row=1, column=1, value=title).font = title_font
            start_row = 3

        # ---- Onglet TABLE ----
        if stype == "table":
            info = tables.get(sdef.get("id"))
            cols = info["columns"]
            rows = info["rows"]
            groups = sdef.get("groups", [])
            ncols = len(cols)
            r = start_row

            if groups:
                # Ligne 1 : groupes fusionnés
                c = 1
                for g in groups:
                    span = max(1, int(g.get("span", 1)))
                    span = min(span, ncols - (c - 1))
                    if span <= 0:
                        break
                    cell = ws.cell(row=r, column=c, value=g.get("label", ""))
                    cell.fill = PatternFill("solid", fgColor=g.get("color", GROUP_FILL))
                    cell.font = group_font
                    cell.alignment = center
                    cell.border = border
                    if span > 1:
                        ws.merge_cells(start_row=r, start_column=c,
                                       end_row=r, end_column=c + span - 1)
                        for cc in range(c + 1, c + span):
                            ws.cell(row=r, column=cc).border = border
                    c += span
                # colonnes restantes sans groupe
                while c <= ncols:
                    ws.cell(row=r, column=c).border = border
                    c += 1
                r += 1

            # Ligne d'en-têtes de colonnes
            for i, col in enumerate(cols, start=1):
                cell = ws.cell(row=r, column=i, value=col.get("label", f"Colonne {i}"))
                cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
                cell.font = header_font
                cell.alignment = center
                cell.border = border
            header_row = r
            r += 1

            # Données typées
            for ri, row in enumerate(rows):
                for i, col in enumerate(cols, start=1):
                    cell = ws.cell(row=r, column=i, value=row[i - 1])
                    cell.border = border
                    nf = _fmt(col)
                    if nf:
                        cell.number_format = nf
                    if ri % 2 == 1:
                        cell.fill = PatternFill("solid", fgColor=ALT_FILL)
                r += 1

            # Ligne de totaux (somme des colonnes numériques) si demandé
            if sdef.get("show_totals") and rows:
                ws.cell(row=r, column=1, value="Total").font = label_font
                for i, col in enumerate(cols, start=1):
                    if col.get("type") in ("integer", "decimal", "currency"):
                        letter = get_column_letter(i)
                        cell = ws.cell(
                            row=r, column=i,
                            value=f"=SUM({letter}{header_row + 1}:{letter}{r - 1})")
                        cell.font = label_font
                        nf = _fmt(col)
                        if nf:
                            cell.number_format = nf
                    ws.cell(row=r, column=i).border = border
                r += 1

            ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
            autosize(ws, ncols)

        # ---- Onglet PIVOT (tableau croisé calculé) ----
        elif stype == "pivot":
            _write_pivot(ws, sdef, tables, pctx, start_row,
                         border, header_font, label_font, center,
                         PatternFill, get_column_letter, autosize)

        # ---- Onglet INFO (cellules fixes typées) ----
        elif stype == "info":
            r = start_row
            for item in sdef.get("items", []):
                label = interpolate(item.get("label", ""), pctx)
                ws.cell(row=r, column=1, value=label).font = label_font
                ws.cell(row=r, column=1).border = border
                val = _coerce(item.get("value", ""), item, pctx)
                cell = ws.cell(row=r, column=2, value=val)
                nf = _fmt(item)
                if nf:
                    cell.number_format = nf
                cell.border = border
                cell.alignment = left
                r += 1
            autosize(ws, 2, minw=16)

    if not wb.sheetnames:
        wb.create_sheet("Feuille 1")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _write_pivot(ws, sdef, tables, pctx, start_row, border, header_font,
                 label_font, center, PatternFill, get_column_letter, autosize):
    """Calcule et écrit un tableau croisé à partir d'un onglet table source."""
    pivot = sdef.get("pivot", {})
    src = tables.get(pivot.get("source"))
    if not src:
        ws.cell(row=start_row, column=1,
                value="Source du tableau croisé introuvable.").font = label_font
        return

    cols = src["columns"]
    rows = src["rows"]
    key_index = {c.get("key", c.get("label")): i for i, c in enumerate(cols)}
    label_index = {c.get("label"): i for i, c in enumerate(cols)}

    def idx(field):
        if field in key_index:
            return key_index[field]
        return label_index.get(field)

    ri = idx(pivot.get("row_field"))
    ci = idx(pivot.get("col_field"))
    vi = idx(pivot.get("value_field"))
    agg = pivot.get("agg", "sum")

    if ri is None or vi is None:
        ws.cell(row=start_row, column=1,
                value="Champs du tableau croisé mal configurés.").font = label_font
        return

    # Catégories
    row_cats, col_cats = [], []
    for row in rows:
        rv = "" if row[ri] is None else str(row[ri])
        if rv not in row_cats:
            row_cats.append(rv)
        if ci is not None:
            cv = "" if row[ci] is None else str(row[ci])
            if cv not in col_cats:
                col_cats.append(cv)
    row_cats.sort()
    col_cats.sort()
    if ci is None:
        col_cats = ["Valeur"]

    # Agrégation
    buckets = {}
    for row in rows:
        rv = "" if row[ri] is None else str(row[ri])
        cv = ("Valeur" if ci is None else
              ("" if row[ci] is None else str(row[ci])))
        num = row[vi]
        if not isinstance(num, (int, float)):
            try:
                num = float(str(num).replace(",", "."))
            except (ValueError, AttributeError):
                num = 0
        buckets.setdefault((rv, cv), []).append(num)

    def aggregate(vals):
        if not vals:
            return None
        if agg == "count":
            return len(vals)
        if agg == "avg":
            return sum(vals) / len(vals)
        if agg == "max":
            return max(vals)
        if agg == "min":
            return min(vals)
        return sum(vals)

    value_fmt = _fmt({"type": pivot.get("value_type", "decimal"),
                      "symbol": pivot.get("symbol", "€"),
                      "decimals": pivot.get("decimals", 2)})

    r0 = start_row
    # Coin + en-têtes de colonnes
    corner = cols[ri].get("label", pivot.get("row_field"))
    hc = ws.cell(row=r0, column=1, value=corner)
    hc.fill = PatternFill("solid", fgColor="1F497D"); hc.font = header_font
    hc.alignment = center; hc.border = border
    for j, cc in enumerate(col_cats, start=2):
        cell = ws.cell(row=r0, column=j, value=cc)
        cell.fill = PatternFill("solid", fgColor="4472C4"); cell.font = header_font
        cell.alignment = center; cell.border = border
    total_col = 2 + len(col_cats)
    tcell = ws.cell(row=r0, column=total_col, value="Total")
    tcell.fill = PatternFill("solid", fgColor="1F497D"); tcell.font = header_font
    tcell.alignment = center; tcell.border = border

    # Lignes
    col_totals = [0] * len(col_cats)
    grand = 0
    for i, rc in enumerate(row_cats, start=1):
        rr = r0 + i
        lc = ws.cell(row=rr, column=1, value=rc); lc.font = label_font
        lc.border = border
        row_total = 0
        for j, cc in enumerate(col_cats):
            val = aggregate(buckets.get((rc, cc), []))
            cell = ws.cell(row=rr, column=2 + j, value=val)
            cell.border = border
            if value_fmt:
                cell.number_format = value_fmt
            if isinstance(val, (int, float)):
                row_total += val
                col_totals[j] += val
        tc = ws.cell(row=rr, column=total_col, value=row_total)
        tc.font = label_font; tc.border = border
        if value_fmt:
            tc.number_format = value_fmt
        grand += row_total

    # Ligne de total général
    rr = r0 + len(row_cats) + 1
    ws.cell(row=rr, column=1, value="Total").font = label_font
    ws.cell(row=rr, column=1).border = border
    for j, ct in enumerate(col_totals):
        cell = ws.cell(row=rr, column=2 + j, value=ct)
        cell.font = label_font; cell.border = border
        if value_fmt:
            cell.number_format = value_fmt
    gc = ws.cell(row=rr, column=total_col, value=grand)
    gc.font = label_font; gc.border = border
    if value_fmt:
        gc.number_format = value_fmt

    ws.freeze_panes = ws.cell(row=r0 + 1, column=2)
    autosize(ws, total_col, minw=12)
