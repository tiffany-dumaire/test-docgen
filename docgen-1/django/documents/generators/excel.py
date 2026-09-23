"""Excel generator.

Template ``config`` shape::

    {
      "sheet_name": "Données",
      "columns": [
        {"key": "region",  "label": "Région",   "group": "Localisation"},
        {"key": "city",    "label": "Ville",     "group": "Localisation"},
        {"key": "product", "label": "Produit",   "group": "Vente"},
        {"key": "qty",     "label": "Quantité",  "group": "Vente", "type": "number"},
        {"key": "amount",  "label": "Montant",   "group": "Vente", "type": "number"}
      ],
      "pivot": {                       # optional cross-tab
        "rows": ["region"],
        "cols": "product",
        "value": "amount",
        "agg": "sum"                   # sum | count | avg
      }
    }

Data rows come from ``document.data['rows']`` (list of dicts) — supplied by
the UI or the external API.
"""
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
GROUP_FILL = PatternFill("solid", fgColor="2E5E8C")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
CENTER = Alignment(horizontal="center", vertical="center")


def _rows(document):
    return (document.data or {}).get("rows", [])


def _agg(values, how):
    nums = [v for v in values if isinstance(v, (int, float))]
    if how == "count":
        return len(values)
    if how == "avg":
        return round(sum(nums) / len(nums), 2) if nums else 0
    return sum(nums)  # default sum


def generate_excel(document, context=None):
    cfg = document.template.config or {}
    columns = cfg.get("columns", [])
    rows = _rows(document)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = cfg.get("sheet_name", "Données")[:31]

    # --- Grouped two-row header (group band + column labels) ---
    has_groups = any(col.get("group") for col in columns)
    start_row = 1
    if has_groups:
        col_idx = 1
        i = 0
        while i < len(columns):
            group = columns[i].get("group", "")
            span = 1
            while i + span < len(columns) and \
                    columns[i + span].get("group", "") == group:
                span += 1
            if group:
                ws.merge_cells(start_row=1, start_column=col_idx,
                               end_row=1, end_column=col_idx + span - 1)
                cell = ws.cell(row=1, column=col_idx, value=group)
                cell.fill = GROUP_FILL
                cell.font = WHITE_BOLD
                cell.alignment = CENTER
            col_idx += span
            i += span
        start_row = 2

    for j, col in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=j, value=col.get("label", col["key"]))
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD
        cell.alignment = CENTER

    # --- Data rows ---
    r = start_row + 1
    for row in rows:
        for j, col in enumerate(columns, start=1):
            ws.cell(row=r, column=j, value=row.get(col["key"]))
        r += 1

    # --- Totals for numeric columns ---
    numeric_cols = [j for j, c in enumerate(columns, start=1)
                    if c.get("type") == "number"]
    if rows and numeric_cols:
        ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
        for j in numeric_cols:
            letter = get_column_letter(j)
            ws.cell(row=r, column=j,
                    value=f"=SUM({letter}{start_row + 1}:{letter}{r - 1})"
                    ).font = Font(bold=True)

    for j in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 18

    # --- Optional pivot / cross-tab sheet ---
    pivot = cfg.get("pivot")
    if pivot and rows:
        _build_pivot(wb, rows, pivot)

    import io
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return f"{document.title}.xlsx", out.read()


def _build_pivot(wb, rows, pivot):
    ps = wb.create_sheet("Tableau croisé")
    row_keys = pivot.get("rows", [])
    col_key = pivot.get("cols")
    value_key = pivot.get("value")
    how = pivot.get("agg", "sum")

    col_values = sorted({str(row.get(col_key, "")) for row in rows})
    buckets = defaultdict(lambda: defaultdict(list))
    for row in rows:
        rkey = tuple(str(row.get(k, "")) for k in row_keys)
        cval = str(row.get(col_key, ""))
        buckets[rkey][cval].append(row.get(value_key))

    # Header
    header = list(row_keys) + col_values + ["Total"]
    for j, label in enumerate(header, start=1):
        cell = ps.cell(row=1, column=j, value=label)
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD
        cell.alignment = CENTER

    r = 2
    for rkey in sorted(buckets):
        for j, part in enumerate(rkey, start=1):
            ps.cell(row=r, column=j, value=part)
        total = 0
        for k, cval in enumerate(col_values, start=len(row_keys) + 1):
            agg = _agg(buckets[rkey][cval], how)
            ps.cell(row=r, column=k, value=agg)
            if isinstance(agg, (int, float)):
                total += agg
        ps.cell(row=r, column=len(header), value=total).font = Font(bold=True)
        r += 1

    for j in range(1, len(header) + 1):
        ps.column_dimensions[get_column_letter(j)].width = 16
