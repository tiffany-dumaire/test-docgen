"""Word generator for analysis documents.

Template ``config`` shape::

    {
      "title": "Analyse {{project.name}}",
      "subtitle": "Client : {{client.name}}",
      "sections": [
        {"heading": "Contexte", "body": "Projet mené pour {{client.name}}..."},
        {"heading": "Résultats", "body": "Chiffre total : {{total}}"}
      ],
      "table": {                       # optional: renders document.data['rows']
        "columns": [
          {"key": "product", "label": "Produit"},
          {"key": "amount",  "label": "Montant"}
        ]
      }
    }
"""
import io

from docx import Document as Docx
from docx.shared import Pt, RGBColor

from documents.context import build_context, resolve


def generate_word(document, context=None):
    cfg = document.template.config or {}
    ctx = build_context(document)
    doc = Docx()

    from core.models import Company
    company = Company.get_solo()
    if company.logo:
        try:
            doc.add_picture(company.logo.path, width=Pt(120))
        except Exception:
            pass

    title = resolve(cfg.get("title", document.title), ctx)
    h = doc.add_heading(title, level=0)

    subtitle = resolve(cfg.get("subtitle", ""), ctx)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.runs[0].italic = True

    # Confidentiality banner
    banner = doc.add_paragraph(
        f"Confidentialité : {document.get_confidentiality_display()}")
    run = banner.runs[0]
    run.bold = True
    run.font.color.rgb = RGBColor(0x8A, 0x2B, 0x2B)

    for section in cfg.get("sections", []):
        doc.add_heading(resolve(section.get("heading", ""), ctx), level=1)
        doc.add_paragraph(resolve(section.get("body", ""), ctx))

    table_cfg = cfg.get("table")
    rows = (document.data or {}).get("rows", [])
    if table_cfg and rows:
        cols = table_cfg.get("columns", [])
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = "Light Grid Accent 1"
        for i, col in enumerate(cols):
            table.rows[0].cells[i].text = col.get("label", col["key"])
        for row in rows:
            cells = table.add_row().cells
            for i, col in enumerate(cols):
                cells[i].text = str(row.get(col["key"], ""))

    doc.add_paragraph()
    footer = doc.add_paragraph(
        f"{company.name} — {company.email} — {company.website_url}")
    footer.runs[0].font.size = Pt(8)

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return f"{document.title}.docx", out.read()
